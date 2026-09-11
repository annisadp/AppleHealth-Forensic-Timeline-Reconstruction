
pip install pandas openpyxl

# **Imports and Cocoa Time Utilities**

import sqlite3
import os
from datetime import datetime, timedelta, timezone
import pandas as pd
from pathlib import Path


COCOA_EPOCH = datetime(
    2001, 1, 1,
    0, 0, 0,
    tzinfo=timezone.utc
)


def parse_cocoa_time(timestamp):

    if timestamp is None:
        return None

    try:
        timestamp = float(timestamp)

        datetime_value = (
            COCOA_EPOCH +
            timedelta(seconds=timestamp)
        )

        return datetime_value.strftime(
            "%Y-%m-%d %H:%M:%S UTC"
        )

    except (
        ValueError,
        TypeError,
        OverflowError
    ):
        return str(timestamp)


print(
    "✅ Libraries imported "
    "and Cocoa Time utilities are ready."
)

# **Mount Google Drive**

from google.colab import drive

drive.mount("/content/drive")

print("✅ Google Drive mounted.")

# **Apple Health Dataset Location**

base_path = (
    "/content/drive/MyDrive/"
    "Colab Notebooks/"
    "Journal Software Impact/"
    "Database iOS Health"
)


folder_mapping = {

    "iOS 13.3.1":
        "iOS_13_3_1",

    "iOS 13.4.1":
        "iOS_13_4_1",

    "iOS 15":
        "iOS_15",

    "iOS 16":
        "iOS_16",

    "iOS 17":
        "iOS_17"
}


print("📁 Dataset root:")
print(base_path)

# **Select iOS Version**

ios_selection = "All Versions (All Datasets)" # @param ["All Versions (All Datasets)", "iOS 13.3.1", "iOS 13.4.1", "iOS 15", "iOS 16", "iOS 17"]


target_datasets = []


if ios_selection == "All Versions (All Datasets)":

    for label, folder in folder_mapping.items():

        target_datasets.append({
            "version": label,
            "path": os.path.join(
                base_path,
                folder
            )
        })

else:

    folder = folder_mapping[ios_selection]

    target_datasets.append({
        "version": ios_selection,
        "path": os.path.join(
            base_path,
            folder
        )
    })


print()
print("=" * 60)
print("🎯 EXTRACTION TARGET")
print("=" * 60)

print(f"iOS Version: {ios_selection}")
print(
    f"Datasets to process: "
    f"{len(target_datasets)}"
)

print("=" * 60)

# **Locate Apple Health Databases**


print()
print("=" * 70)
print("🔍 APPLE HEALTH DATABASE VERIFICATION")
print("=" * 70)


for dataset in target_datasets:

    version = dataset["version"]
    folder_path = dataset["path"]


    print()
    print(f"📱 {version}")
    print(f"📁 {folder_path}")


    if not os.path.exists(folder_path):

        print("   ❌ Dataset folder not found.")

        dataset["healthdb"] = None
        dataset["healthdb_secure"] = None

        continue


    files = os.listdir(folder_path)


    healthdb_path = None
    healthdb_secure_path = None


    for filename in files:

        filename_lower = filename.lower()


        if filename_lower.endswith(
            "healthdb_secure.sqlite"
        ):

            healthdb_secure_path = os.path.join(
                folder_path,
                filename
            )


        elif filename_lower.endswith(
            "healthdb.sqlite"
        ):

            healthdb_path = os.path.join(
                folder_path,
                filename
            )


    dataset["healthdb"] = healthdb_path

    dataset["healthdb_secure"] = (
        healthdb_secure_path
    )


    if healthdb_path:

        print(
            "   ├─ healthdb.sqlite"
            "        : ✅ Found"
        )

        print(
            f"   │  └─ {os.path.basename(healthdb_path)}"
        )

    else:

        print(
            "   ├─ healthdb.sqlite"
            "        : ❌ Not found"
        )


    if healthdb_secure_path:

        print(
            "   └─ healthdb_secure.sqlite"
            " : ✅ Found"
        )

        print(
            f"      └─ "
            f"{os.path.basename(healthdb_secure_path)}"
        )

    else:

        print(
            "   └─ healthdb_secure.sqlite"
            " : ❌ Not found"
        )


print()
print("=" * 70)

"""# **Dataset Readiness Summary**"""

print()
print("=" * 70)
print("📋 DATASET READINESS")
print("=" * 70)


ready_dataset_count = 0


for dataset in target_datasets:

    version = dataset["version"]

    healthdb = dataset.get("healthdb")

    healthdb_secure = dataset.get(
        "healthdb_secure"
    )


    if healthdb and healthdb_secure:

        status = "✅ READY"

        ready_dataset_count += 1

    else:

        status = "❌ INCOMPLETE"


    print(
        f"{version:<12} : {status}"
    )


print("-" * 70)

print(
    f"Ready datasets: "
    f"{ready_dataset_count}/"
    f"{len(target_datasets)}"
)

print("=" * 70)

# **Read-Only SQLite Connection**

def open_sqlite_readonly(database_path):

    db_path = Path(database_path)

    db_uri = db_path.resolve().as_uri() + "?mode=ro"

    conn = sqlite3.connect(
        db_uri,
        uri=True
    )

    conn.row_factory = sqlite3.Row

    conn.execute("PRAGMA query_only = ON")

    return conn


print("✅ Read-only SQLite connection is ready.")

"""# **Validate SQLite Databases**"""

print()
print("=" * 70)
print("🔍 SQLITE DATABASE VALIDATION")
print("=" * 70)


for dataset in target_datasets:

    version = dataset["version"]

    healthdb_path = dataset.get("healthdb")
    healthdb_secure_path = dataset.get(
        "healthdb_secure"
    )

    print()
    print(f"📱 {version}")


    if healthdb_path:

        try:

            conn = open_sqlite_readonly(
                healthdb_path
            )

            result = conn.execute(
                "PRAGMA quick_check;"
            ).fetchone()

            conn.close()

            print(
                f"   ├─ healthdb.sqlite        : "
                f"✅ Valid SQLite ({result[0]})"
            )

        except Exception as e:

            print(
                f"   ├─ healthdb.sqlite        : "
                f"❌ Error"
            )

            print(
                f"   │  └─ {e}"
            )


    if healthdb_secure_path:

        try:

            conn = open_sqlite_readonly(
                healthdb_secure_path
            )

            result = conn.execute(
                "PRAGMA quick_check;"
            ).fetchone()

            conn.close()

            print(
                f"   └─ healthdb_secure.sqlite : "
                f"✅ Valid SQLite ({result[0]})"
            )

        except Exception as e:

            print(
                f"   └─ healthdb_secure.sqlite : "
                f"❌ Error"
            )

            print(
                f"      └─ {e}"
            )


print()
print("=" * 70)

# **Inspect Database Tables**

def get_sqlite_tables(database_path):

    conn = open_sqlite_readonly(
        database_path
    )

    query = """
    SELECT name
    FROM sqlite_master
    WHERE type = 'table'
    ORDER BY name;
    """

    rows = conn.execute(query).fetchall()

    tables = [
        row["name"]
        for row in rows
    ]

    conn.close()

    return tables


print("✅ Table inspection is ready.")

print()
print("=" * 70)
print("📊 APPLE HEALTH DATABASE STRUCTURE")
print("=" * 70)


for dataset in target_datasets:

    version = dataset["version"]

    healthdb_path = dataset["healthdb"]

    healthdb_secure_path = dataset[
        "healthdb_secure"
    ]


    print()
    print(f"📱 {version}")
    print("-" * 70)


    healthdb_tables = get_sqlite_tables(
        healthdb_path
    )

    print(
        f"📂 healthdb.sqlite"
        f" → {len(healthdb_tables)} tables"
    )


    secure_tables = get_sqlite_tables(
        healthdb_secure_path
    )

    print(
        f"🔐 healthdb_secure.sqlite"
        f" → {len(secure_tables)} tables"
    )


    dataset["healthdb_tables"] = (
        healthdb_tables
    )

    dataset["healthdb_secure_tables"] = (
        secure_tables
    )


print()
print("=" * 70)

for dataset in target_datasets:

    print()
    print("=" * 70)
    print(f"📱 {dataset['version']}")
    print("=" * 70)


    print()
    print("📂 healthdb.sqlite TABLES")
    print("-" * 70)

    for table in dataset[
        "healthdb_tables"
    ]:

        print(table)


    print()
    print("🔐 healthdb_secure.sqlite TABLES")
    print("-" * 70)

    for table in dataset[
        "healthdb_secure_tables"
    ]:

        print(table)

# **Achievements**


ACHIEVEMENT_TABLE = (
    "ACHAchievementsPlugin_earned_instances"
)


print()
print("=" * 70)
print("🏆 ACHIEVEMENTS TABLE CHECK")
print("=" * 70)


for dataset in target_datasets:

    version = dataset["version"]

    secure_tables = dataset[
        "healthdb_secure_tables"
    ]


    if ACHIEVEMENT_TABLE in secure_tables:

        print(
            f"{version:<12} : "
            f"✅ Table found"
        )

    else:

        print(
            f"{version:<12} : "
            f"❌ Table not found"
        )


print("=" * 70)


def get_table_columns(
    database_path,
    table_name
):

    conn = open_sqlite_readonly(
        database_path
    )

    query = (
        f'PRAGMA table_info("{table_name}");'
    )

    rows = conn.execute(
        query
    ).fetchall()

    conn.close()


    columns = []

    for row in rows:

        columns.append({
            "cid": row["cid"],
            "name": row["name"],
            "type": row["type"],
            "notnull": row["notnull"],
            "pk": row["pk"]
        })


    return columns


for dataset in target_datasets:

    version = dataset["version"]

    secure_path = dataset[
        "healthdb_secure"
    ]

    secure_tables = dataset[
        "healthdb_secure_tables"
    ]


    print()
    print("=" * 70)
    print(f"🏆 ACHIEVEMENTS - {version}")
    print("=" * 70)


    if ACHIEVEMENT_TABLE not in secure_tables:

        print(
            "❌ Achievements table "
            "not found."
        )

        continue


    columns = get_table_columns(
        secure_path,
        ACHIEVEMENT_TABLE
    )


    df_columns = pd.DataFrame(
        columns
    )


    display(df_columns)


def extract_achievements(database_path):

    conn = open_sqlite_readonly(database_path)

    query = """
    SELECT
        ROWID,
        template_unique_name,
        created_date,
        earned_date,
        value_in_canonical_unit,
        value_canonical_unit,
        creator_device,
        sync_provenance

    FROM
        ACHAchievementsPlugin_earned_instances
    """

    df = pd.read_sql_query(
        query,
        conn
    )

    conn.close()

    return df


print("✅ Achievements extractor is ready.")


achievement_results = {}


for dataset in target_datasets:

    version = dataset["version"]

    secure_path = dataset[
        "healthdb_secure"
    ]

    print()
    print("=" * 70)
    print(f"🏆 ACHIEVEMENTS EXTRACTION - {version}")
    print("=" * 70)


    try:

        df_achievements = extract_achievements(
            secure_path
        )


        achievement_results[
            version
        ] = df_achievements


        print(
            f"✅ Extraction successful"
        )

        print(
            f"📊 Total events: "
            f"{len(df_achievements)}"
        )


    except Exception as e:

        print(
            f"❌ Extraction failed: {e}"
        )


for version, df in achievement_results.items():

    print()
    print("=" * 70)
    print(f"🏆 RAW ACHIEVEMENTS DATA - {version}")
    print("=" * 70)

    display(
        df.head(10)
    )


def parse_cocoa_datetime(timestamp):

    if timestamp is None:
        return pd.NaT

    try:
        timestamp = float(timestamp)

        dt = (
            COCOA_EPOCH
            + timedelta(seconds=timestamp)
        )

        return dt

    except (
        ValueError,
        TypeError,
        OverflowError
    ):
        return pd.NaT


achievement_reports = {}


for version, df in achievement_results.items():

    df_report = pd.DataFrame()

    df_report["Created Timestamp"] = (
        df["created_date"]
        .apply(parse_cocoa_datetime)
    )

    df_report["Earned Date"] = (
        df["earned_date"]
    )

    df_report["Achievement"] = (
        df["template_unique_name"]
    )

    df_report["Value"] = (
        df["value_in_canonical_unit"]
    )

    df_report["Unit"] = (
        df["value_canonical_unit"]
    )

    df_report["Creator Device"] = (
        df["creator_device"]
    )

    df_report["Sync Provenance"] = (
        df["sync_provenance"]
    )

    df_report["Source"] = (
        "healthdb_secure.sqlite"
    )

    achievement_reports[
        version
    ] = df_report


print(
    "✅ Achievements report "
    "created."
)


for version, df_report in achievement_reports.items():

    print()
    print("=" * 70)
    print(f"🏆 ACHIEVEMENTS - {version}")
    print("=" * 70)

    display(
        df_report.head(20)
    )

    print()
    print(
        f"📊 Total Achievements Events: "
        f"{len(df_report)}"
    )

# **All Watch Sleep**


def seconds_to_hms(seconds_value):

    if seconds_value is None:
        return None

    try:
        total = int(float(seconds_value))

    except (TypeError, ValueError):
        return None


    hours = total // 3600

    minutes = (
        total % 3600
    ) // 60

    seconds = total % 60


    return (
        f"{hours:02d}:"
        f"{minutes:02d}:"
        f"{seconds:02d}"
    )


print(
    "✅ Duration conversion is ready."
)


def extract_all_watch_sleep(
    database_path,
    ios_version
):

    conn = open_sqlite_readonly(
        database_path
    )


    if ios_version == "iOS 17":

        query = """
        SELECT
            s.start_date AS start_cocoa,
            s.end_date AS end_cocoa,
            cs.value AS state_code

        FROM samples s

        JOIN category_samples cs
            ON s.data_id = cs.data_id

        ORDER BY
            s.start_date ASC
        """


        df = pd.read_sql_query(
            query,
            conn
        )


        df["state_code"] = pd.to_numeric(
            df["state_code"],
            errors="coerce"
        )


        df = df[
            df["state_code"].isin(
                [2, 3, 4, 5]
            )
        ].copy()


        df["state_code"] = (
            df["state_code"]
            .astype(int)
        )


    else:

        query = """
        SELECT
            s.start_date AS start_cocoa,
            s.end_date AS end_cocoa,
            cs.value AS state_code

        FROM samples s

        LEFT JOIN category_samples cs
            ON s.data_id = cs.data_id

        LEFT JOIN objects o
            ON s.data_id = o.data_id

        LEFT JOIN data_provenances dp
            ON o.provenance = dp.rowid
            AND dp.origin_product_type LIKE '%Watch%'

        WHERE
            s.data_type IN (63, '63')
            AND cs.value IN (0, 1)

        ORDER BY
            s.start_date ASC
        """


        df = pd.read_sql_query(
            query,
            conn
        )


    conn.close()


    return df


print(
    "✅ All Watch Sleep extractor is ready."
)


all_watch_sleep_results = {}


for dataset in target_datasets:

    version = dataset[
        "version"
    ]

    secure_path = dataset[
        "healthdb_secure"
    ]


    print()
    print("=" * 70)

    print(
        f"😴 ALL WATCH SLEEP - {version}"
    )

    print("=" * 70)


    try:

        df_sleep = (
            extract_all_watch_sleep(
                secure_path,
                version
            )
        )


        all_watch_sleep_results[
            version
        ] = df_sleep


        print(
            "✅ Extraction successful"
        )

        print(
            f"📊 Total Events: "
            f"{len(df_sleep)}"
        )


        if version == "iOS 17":

            print(
                "📱 Schema: "
                "iOS 17"
            )

            print(
                "💤 Sleep State: "
                "2, 3, 4, 5"
            )

        else:

            print(
                "📱 Schema: "
                "iOS 13–16"
            )

            print(
                "💤 Sleep State: "
                "0, 1"
            )


    except Exception as e:

        print(
            f"❌ Extraction failed: {e}"
        )


for version, df_sleep in all_watch_sleep_results.items():

    print()
    print("=" * 70)

    print(
        f"😴 RAW ALL WATCH SLEEP - {version}"
    )

    print("=" * 70)


    display(
        df_sleep.head(20)
    )


    print()

    print(
        f"📊 Total Events: "
        f"{len(df_sleep)}"
    )


SLEEP_STATE_MAP = {
    0: "In Bed",
    1: "Asleep",
    2: "Awake",
    3: "Core",
    4: "Deep",
    5: "REM"
}


all_watch_sleep_reports = {}


for version, df_sleep in all_watch_sleep_results.items():

    df_report = pd.DataFrame()


    df_report["Start Timestamp"] = (
        df_sleep["start_cocoa"]
        .apply(parse_cocoa_datetime)
    )


    df_report["End Timestamp"] = (
        df_sleep["end_cocoa"]
        .apply(parse_cocoa_datetime)
    )


    df_report["Sleep State Code"] = (
        df_sleep["state_code"]
    )


    df_report["Sleep State"] = (
        df_sleep["state_code"]
        .map(SLEEP_STATE_MAP)
    )


    duration_seconds = (
        df_sleep["end_cocoa"]
        -
        df_sleep["start_cocoa"]
    )


    df_report["Sleep Duration"] = (
        duration_seconds
        .apply(seconds_to_hms)
    )


    df_report["Source"] = (
        "healthdb_secure.sqlite"
    )


    all_watch_sleep_reports[
        version
    ] = df_report


print(
    "✅ All Watch Sleep report "
    "created."
)


for version, df_report in all_watch_sleep_reports.items():

    print()
    print("=" * 70)

    print(
        f"😴 ALL WATCH SLEEP REPORT - {version}"
    )

    print("=" * 70)


    display(
        df_report.head(20)
    )


    print()

    print(
        f"📊 Total All Watch Sleep Events: "
        f"{len(df_report)}"
    )

# **Headphone Audio Levels**


HEADPHONE_SECURE_TABLES = [
    "samples",
    "quantity_samples",
    "metadata_values",
    "metadata_keys",
    "objects",
    "data_provenances"
]

HEADPHONE_HEALTHDB_TABLES = [
    "source_devices"
]


print()
print("=" * 70)
print("🎧 VERIFIKASI HEADPHONE AUDIO LEVELS")
print("=" * 70)


for dataset in target_datasets:

    version = dataset["version"]

    secure_tables = dataset[
        "healthdb_secure_tables"
    ]

    healthdb_tables = dataset[
        "healthdb_tables"
    ]


    print()
    print(f"📱 {version}")


    print("   🔐 healthdb_secure.sqlite")

    for table in HEADPHONE_SECURE_TABLES:

        status = (
            "✅"
            if table in secure_tables
            else "❌"
        )

        print(
            f"      {status} {table}"
        )


    print("   📂 healthdb.sqlite")

    for table in HEADPHONE_HEALTHDB_TABLES:

        status = (
            "✅"
            if table in healthdb_tables
            else "❌"
        )

        print(
            f"      {status} {table}"
        )


print()
print("=" * 70)


def extract_headphone_audio_levels(
    secure_database_path,
    health_database_path
):


    secure_conn = open_sqlite_readonly(
        secure_database_path
    )


    secure_query = """
    SELECT

        samples.start_date
            AS start_cocoa,

        samples.end_date
            AS end_cocoa,

        quantity_samples.quantity
            AS decibels,

        metadata_values.string_value
            AS bundle_name,

        metadata_keys.key
            AS key,

        samples.data_id
            AS data_id,

        data_provenances.device_id
            AS device_id

    FROM samples


    LEFT JOIN quantity_samples

        ON quantity_samples.data_id
        = samples.data_id


    LEFT JOIN metadata_values

        ON metadata_values.object_id
        = samples.data_id


    LEFT JOIN metadata_keys

        ON metadata_keys.ROWID
        = metadata_values.key_id


    LEFT JOIN objects

        ON objects.data_id
        = samples.data_id


    LEFT JOIN data_provenances

        ON data_provenances.ROWID
        = objects.provenance


    WHERE

        samples.data_type IN (173, '173')

        AND (

            metadata_keys.key IS NULL

            OR metadata_keys.key !=
            '_HKPrivateMetadataKeyHeadphoneAudioDataIsTransient'

        )


    GROUP BY
        samples.data_id


    ORDER BY
        samples.start_date
    """


    df_secure = pd.read_sql_query(
        secure_query,
        secure_conn
    )


    secure_conn.close()


    health_conn = open_sqlite_readonly(
        health_database_path
    )


    device_query = """
    SELECT

        ROWID
            AS device_id,

        name
            AS device_name,

        manufacturer
            AS device_manufacturer,

        model
            AS device_model,

        localIdentifier
            AS local_identifier

    FROM source_devices
    """


    df_devices = pd.read_sql_query(
        device_query,
        health_conn
    )


    health_conn.close()


    df_secure["device_id"] = pd.to_numeric(
        df_secure["device_id"],
        errors="coerce"
    ).astype("Int64")


    df_devices["device_id"] = pd.to_numeric(
        df_devices["device_id"],
        errors="coerce"
    ).astype("Int64")


    df_result = df_secure.merge(

        df_devices,

        how="left",

        on="device_id",

        validate="many_to_one"
    )


    return df_result


print(
    "✅ Extractor Headphone Audio Levels "
    "cross-database extractor is ready."
)


headphone_audio_results = {}


for dataset in target_datasets:

    version = dataset["version"]

    secure_path = dataset[
        "healthdb_secure"
    ]

    healthdb_path = dataset[
        "healthdb"
    ]


    print()
    print("=" * 70)
    print(
        f"🎧 HEADPHONE AUDIO LEVELS - {version}"
    )
    print("=" * 70)


    try:

        df_audio = (
            extract_headphone_audio_levels(
                secure_path,
                healthdb_path
            )
        )


        headphone_audio_results[
            version
        ] = df_audio


        print(
            "✅ Extraction successful"
        )

        print(
            f"📊 Total Events: "
            f"{len(df_audio)}"
        )


        device_found = (
            df_audio[
                [
                    "device_name",
                    "device_manufacturer",
                    "device_model",
                    "local_identifier"
                ]
            ]
            .notna()
            .any(axis=1)
            .sum()
        )


        print(
            f"🔗 Events with device metadata: "
            f"{device_found}"
        )


        print(
            "🔐 Primary data  : "
            "healthdb_secure.sqlite"
        )

        print(
            "📂 Device data : "
            "healthdb.sqlite"
        )


    except Exception as e:

        print(
            f"❌ Extraction failed: {e}"
        )


for version, df_audio in headphone_audio_results.items():

    print()
    print("=" * 70)

    print(
        f"🎧 RAW HEADPHONE AUDIO LEVELS - {version}"
    )

    print("=" * 70)


    display(
        df_audio.head(20)
    )


    print()

    print(
        f"📊 Total Events: "
        f"{len(df_audio)}"
    )


headphone_audio_reports = {}


for version, df_audio in headphone_audio_results.items():

    df_report = pd.DataFrame()


    df_report[
        "Start Timestamp"
    ] = (
        df_audio["start_cocoa"]
        .apply(parse_cocoa_datetime)
    )


    df_report[
        "End Timestamp"
    ] = (
        df_audio["end_cocoa"]
        .apply(parse_cocoa_datetime)
    )


    duration_seconds = (

        df_audio["end_cocoa"]

        -

        df_audio["start_cocoa"]
    )


    df_report[
        "Total Time Duration"
    ] = (
        duration_seconds
        .apply(seconds_to_hms)
    )


    df_report[
        "Decibels"
    ] = (
        df_audio["decibels"]
    )


    df_report[
        "Bundle Name"
    ] = (
        df_audio["bundle_name"]
    )


    df_report[
        "Device Name"
    ] = (
        df_audio["device_name"]
    )


    df_report[
        "Device Manufacturer"
    ] = (
        df_audio["device_manufacturer"]
    )


    df_report[
        "Device Model"
    ] = (
        df_audio["device_model"]
    )


    df_report[
        "Local Identifier"
    ] = (
        df_audio["local_identifier"]
    )


    df_report[
        "Key"
    ] = (
        df_audio["key"]
    )


    df_report[
        "Data ID"
    ] = (
        df_audio["data_id"]
    )


    df_report[
        "Health Data Source"
    ] = (
        "healthdb_secure.sqlite"
    )


    df_report[
        "Device Metadata Source"
    ] = (
        "healthdb.sqlite"
    )


    headphone_audio_reports[
        version
    ] = df_report


print(
    "✅ Headphone Audio Levels report "
    "created."
)


for version, df_report in headphone_audio_reports.items():

    print()
    print("=" * 70)

    print(
        f"🎧 HEADPHONE AUDIO LEVELS REPORT - {version}"
    )

    print("=" * 70)


    display(
        df_report.head(20)
    )


    print()

    print(
        f"📊 Total Headphone Audio Events: "
        f"{len(df_report)}"
    )

# **Heart Rate**


HEART_RATE_SECURE_TABLES = [
    "samples",
    "quantity_samples",
    "metadata_values",
    "objects",
    "data_provenances"
]

HEART_RATE_HEALTHDB_TABLES = [
    "source_devices",
    "sources"
]


print()
print("=" * 70)
print("❤️ HEART RATE STRUCTURE VERIFICATION")
print("=" * 70)


for dataset in target_datasets:

    version = dataset["version"]

    secure_tables = dataset[
        "healthdb_secure_tables"
    ]

    healthdb_tables = dataset[
        "healthdb_tables"
    ]


    print()
    print(f"📱 {version}")


    print("   🔐 healthdb_secure.sqlite")

    for table in HEART_RATE_SECURE_TABLES:

        status = (
            "✅"
            if table in secure_tables
            else "❌"
        )

        print(
            f"      {status} {table}"
        )


    print("   📂 healthdb.sqlite")

    for table in HEART_RATE_HEALTHDB_TABLES:

        status = (
            "✅"
            if table in healthdb_tables
            else "❌"
        )

        print(
            f"      {status} {table}"
        )


print()
print("=" * 70)


def extract_heart_rate(
    secure_database_path,
    health_database_path
):


    secure_conn = open_sqlite_readonly(
        secure_database_path
    )


    secure_query = """
    SELECT

        s.start_date
            AS start_cocoa,

        s.end_date
            AS end_cocoa,

        CAST(
            ROUND(qs.quantity * 60.0)
            AS INTEGER
        )
            AS bpm,

        mv.numerical_value
            AS context_code,

        o.creation_date
            AS added_cocoa,

        dp.device_id
            AS device_id,

        dp.source_id
            AS source_id,

        dp.source_version
            AS software_version,

        dp.tz_name
            AS tz_name,

        s.data_id
            AS data_id

    FROM samples s


    LEFT JOIN quantity_samples qs

        ON qs.data_id
        = s.data_id


    LEFT JOIN metadata_values mv

        ON mv.object_id
        = s.data_id


    LEFT JOIN objects o

        ON o.data_id
        = s.data_id


    LEFT JOIN data_provenances dp

        ON dp.ROWID
        = o.provenance


    WHERE

        s.data_type IN (5, '5')

        AND o.type != 2


    ORDER BY

        s.start_date DESC
    """


    df_secure = pd.read_sql_query(
        secure_query,
        secure_conn
    )


    secure_conn.close()


    health_conn = open_sqlite_readonly(
        health_database_path
    )


    device_query = """
    SELECT

        ROWID
            AS device_id,

        name
            AS device_name,

        manufacturer
            AS manufacturer,

        hardware
            AS hardware

    FROM source_devices
    """


    df_devices = pd.read_sql_query(
        device_query,
        health_conn
    )


    source_query = """
    SELECT

        ROWID
            AS source_id,

        name
            AS source_name,

        source_options
            AS source_options

    FROM sources
    """


    df_sources = pd.read_sql_query(
        source_query,
        health_conn
    )


    health_conn.close()


    for dataframe, column in [

        (df_secure, "device_id"),
        (df_devices, "device_id"),
        (df_secure, "source_id"),
        (df_sources, "source_id")

    ]:

        dataframe[column] = pd.to_numeric(
            dataframe[column],
            errors="coerce"
        ).astype("Int64")


    df_result = df_secure.merge(

        df_devices,

        how="left",

        on="device_id",

        validate="many_to_one"
    )


    df_result = df_result.merge(

        df_sources,

        how="left",

        on="source_id",

        validate="many_to_one"
    )


    df_result["device_name"] = (
        df_result["device_name"]
        .replace("**NONE**", "")
    )


    return df_result


print(
    "✅ Extractor Heart Rate "
    "cross-database extractor is ready."
)


heart_rate_results = {}


for dataset in target_datasets:

    version = dataset["version"]

    secure_path = dataset[
        "healthdb_secure"
    ]

    healthdb_path = dataset[
        "healthdb"
    ]


    print()
    print("=" * 70)

    print(
        f"❤️ HEART RATE - {version}"
    )

    print("=" * 70)


    try:

        df_heart_rate = extract_heart_rate(
            secure_path,
            healthdb_path
        )


        heart_rate_results[
            version
        ] = df_heart_rate


        print(
            "✅ Extraction successful"
        )


        print(
            f"📊 Total Events: "
            f"{len(df_heart_rate)}"
        )


        device_found = (
            df_heart_rate[
                "device_name"
            ]
            .notna()
            .sum()
        )


        source_found = (
            df_heart_rate[
                "source_name"
            ]
            .notna()
            .sum()
        )


        print(
            f"⌚ Events with Device : "
            f"{device_found}"
        )


        print(
            f"📱 Events with Source : "
            f"{source_found}"
        )


        print(
            "🔐 Health Data  : "
            "healthdb_secure.sqlite"
        )


        print(
            "📂 Device/Source: "
            "healthdb.sqlite"
        )


    except Exception as e:

        print(
            f"❌ Extraction failed: {e}"
        )


for version, df_heart_rate in heart_rate_results.items():

    print()
    print("=" * 70)

    print(
        f"❤️ RAW HEART RATE - {version}"
    )

    print("=" * 70)


    display(
        df_heart_rate.head(20)
    )


    print()

    print(
        f"📊 Total Events: "
        f"{len(df_heart_rate)}"
    )


HEART_RATE_CONTEXT_MAP = {

    1: "Background",
    2: "Streaming",
    3: "Sedentary",
    4: "Walking",
    5: "Breathe",
    6: "Workout",
    8: "Background",
    9: "ECG",
    10: "Blood Oxygen Saturation"

}


def map_heart_rate_context(value):

    if pd.isna(value):
        return None

    try:

        numeric_value = float(value)

        int_value = int(
            numeric_value
        )

        if numeric_value == int_value:

            if int_value in HEART_RATE_CONTEXT_MAP:

                return HEART_RATE_CONTEXT_MAP[
                    int_value
                ]

        return str(value)

    except (
        TypeError,
        ValueError
    ):

        return str(value)


print(
    "✅ Heart Rate context mapping is ready."
)


heart_rate_reports = {}


for version, df_heart_rate in heart_rate_results.items():

    df_report = pd.DataFrame()


    df_report[
        "Start Timestamp"
    ] = (
        df_heart_rate[
            "start_cocoa"
        ]
        .apply(
            parse_cocoa_datetime
        )
    )


    df_report[
        "End Timestamp"
    ] = (
        df_heart_rate[
            "end_cocoa"
        ]
        .apply(
            parse_cocoa_datetime
        )
    )


    df_report[
        "Added Timestamp"
    ] = (
        df_heart_rate[
            "added_cocoa"
        ]
        .apply(
            parse_cocoa_datetime
        )
    )


    df_report[
        "BPM"
    ] = (
        df_heart_rate[
            "bpm"
        ]
    )


    df_report[
        "Context"
    ] = (
        df_heart_rate[
            "context_code"
        ]
        .apply(
            map_heart_rate_context
        )
    )


    df_report[
        "Device Name"
    ] = (
        df_heart_rate[
            "device_name"
        ]
    )


    df_report[
        "Manufacturer"
    ] = (
        df_heart_rate[
            "manufacturer"
        ]
    )


    df_report[
        "Hardware"
    ] = (
        df_heart_rate[
            "hardware"
        ]
    )


    df_report[
        "Source Name"
    ] = (
        df_heart_rate[
            "source_name"
        ]
    )


    df_report[
        "Software Version"
    ] = (
        df_heart_rate[
            "software_version"
        ]
    )


    df_report[
        "Time Zone"
    ] = (
        df_heart_rate[
            "tz_name"
        ]
    )


    df_report[
        "Source Options"
    ] = (
        df_heart_rate[
            "source_options"
        ]
    )


    df_report[
        "Data ID"
    ] = (
        df_heart_rate[
            "data_id"
        ]
    )


    df_report[
        "Health Data Source"
    ] = (
        "healthdb_secure.sqlite"
    )


    df_report[
        "Device & Source Metadata"
    ] = (
        "healthdb.sqlite"
    )


    heart_rate_reports[
        version
    ] = df_report


print(
    "✅ Heart Rate report "
    "created."
)


for version, df_report in heart_rate_reports.items():

    print()
    print("=" * 70)

    print(
        f"❤️ HEART RATE REPORT - {version}"
    )

    print("=" * 70)


    display(
        df_report.head(20)
    )


    print()

    print(
        f"📊 Total Heart Rate Events: "
        f"{len(df_report)}"
    )

"""# **Height**"""


HEIGHT_REQUIRED_TABLES = [
    "samples",
    "quantity_samples"
]


print()
print("=" * 70)
print("📏 HEIGHT STRUCTURE VERIFICATION")
print("=" * 70)


for dataset in target_datasets:

    version = dataset["version"]

    secure_tables = dataset[
        "healthdb_secure_tables"
    ]


    print()
    print(f"📱 {version}")

    for table in HEIGHT_REQUIRED_TABLES:

        status = (
            "✅"
            if table in secure_tables
            else "❌"
        )

        print(
            f"   🔐 {status} {table}"
        )


print()
print("=" * 70)


def extract_height(database_path):


    conn = open_sqlite_readonly(
        database_path
    )


    query = """
    SELECT

        samples.start_date
            AS start_cocoa,

        quantity_samples.quantity
            AS height_meters

    FROM samples

    LEFT JOIN quantity_samples

        ON samples.data_id
        = quantity_samples.data_id

    WHERE

        samples.data_type IN (2, '2')

        AND quantity_samples.quantity IS NOT NULL

    ORDER BY
        samples.start_date DESC
    """


    df = pd.read_sql_query(
        query,
        conn
    )


    conn.close()


    return df


print(
    "✅ Height extractor is ready."
)


height_results = {}


for dataset in target_datasets:

    version = dataset["version"]

    secure_path = dataset[
        "healthdb_secure"
    ]


    print()
    print("=" * 70)

    print(
        f"📏 HEIGHT - {version}"
    )

    print("=" * 70)


    try:

        df_height = extract_height(
            secure_path
        )


        height_results[
            version
        ] = df_height


        print(
            "✅ Extraction successful"
        )


        print(
            f"📊 Total Events: "
            f"{len(df_height)}"
        )


        print(
            "🔐 Source: "
            "healthdb_secure.sqlite"
        )


    except Exception as e:

        print(
            f"❌ Extraction failed: {e}"
        )


for version, df_height in height_results.items():

    print()
    print("=" * 70)

    print(
        f"📏 RAW HEIGHT - {version}"
    )

    print("=" * 70)


    display(
        df_height.head(20)
    )


    print()

    print(
        f"📊 Total Events: "
        f"{len(df_height)}"
    )


def meters_to_centimeters_int(meters):


    try:

        return int(
            float(meters) * 100
        )

    except (
        TypeError,
        ValueError
    ):

        return None


def meters_to_feet_inches(meters):


    if meters is None:

        return None


    try:

        feet_float = (
            float(meters)
            * 3.280839895
        )


        feet = int(
            feet_float
        )


        inches = int(
            (
                (feet_float - feet)
                * 12
            )
            + 0.5
        )


        if inches == 12:

            feet += 1

            inches = 0


        return (
            f"{feet}'{inches}\""
        )


    except (
        TypeError,
        ValueError
    ):

        return None


print(
    "✅ Height conversion is ready."
)


height_reports = {}


for version, df_height in height_results.items():

    df_report = pd.DataFrame()


    df_report[
        "Timestamp"
    ] = (
        df_height[
            "start_cocoa"
        ]
        .apply(
            parse_cocoa_datetime
        )
    )


    df_report[
        "Height (m)"
    ] = (
        df_height[
            "height_meters"
        ]
    )


    df_report[
        "Height (cm)"
    ] = (
        df_height[
            "height_meters"
        ]
        .apply(
            meters_to_centimeters_int
        )
    )


    df_report[
        "Height (ft/in)"
    ] = (
        df_height[
            "height_meters"
        ]
        .apply(
            meters_to_feet_inches
        )
    )


    df_report[
        "Source"
    ] = (
        "healthdb_secure.sqlite"
    )


    height_reports[
        version
    ] = df_report


print(
    "✅ Height report formatted."
)


for version, df_report in height_reports.items():

    print()
    print("=" * 70)

    print(
        f"📏 HEIGHT REPORT - {version}"
    )

    print("=" * 70)


    display(
        df_report.head(20)
    )


    print()

    print(
        f"📊 Total Height Events: "
        f"{len(df_report)}"
    )

"""# **Resting Heart Rate**"""


RESTING_HR_SECURE_TABLES = [
    "samples",
    "quantity_samples",
    "objects",
    "data_provenances"
]


RESTING_HR_HEALTHDB_TABLES = [
    "sources"
]


print()
print("=" * 70)
print("❤️‍🩹 RESTING HEART RATE STRUCTURE VERIFICATION")
print("=" * 70)


for dataset in target_datasets:

    version = dataset["version"]


    secure_tables = dataset[
        "healthdb_secure_tables"
    ]


    healthdb_tables = dataset[
        "healthdb_tables"
    ]


    print()
    print(f"📱 {version}")


    print(
        "   🔐 healthdb_secure.sqlite"
    )


    for table in RESTING_HR_SECURE_TABLES:

        status = (
            "✅"
            if table in secure_tables
            else "❌"
        )

        print(
            f"      {status} {table}"
        )


    print(
        "   📂 healthdb.sqlite"
    )


    for table in RESTING_HR_HEALTHDB_TABLES:

        status = (
            "✅"
            if table in healthdb_tables
            else "❌"
        )


        print(
            f"      {status} {table}"
        )


print()
print("=" * 70)


def extract_resting_heart_rate(
    secure_database_path,
    health_database_path
):


    secure_conn = open_sqlite_readonly(
        secure_database_path
    )


    secure_query = """

    SELECT


        samples.start_date
            AS start_cocoa,


        samples.end_date
            AS end_cocoa,


        CAST(
            ROUND(quantity_samples.quantity)
            AS INTEGER
        )
            AS resting_heart_rate,


        objects.creation_date
            AS added_cocoa,


        data_provenances.source_id
            AS source_id,


        samples.data_id
            AS data_id


    FROM samples


    LEFT JOIN quantity_samples

        ON samples.data_id =
           quantity_samples.data_id


    LEFT JOIN objects

        ON samples.data_id =
           objects.data_id


    LEFT JOIN data_provenances

        ON objects.provenance =
           data_provenances.ROWID


    WHERE


        samples.data_type IN (118, '118')


        AND quantity_samples.quantity IS NOT NULL


    ORDER BY

        samples.start_date DESC

    """


    df_secure = pd.read_sql_query(
        secure_query,
        secure_conn
    )


    secure_conn.close()


    health_conn = open_sqlite_readonly(
        health_database_path
    )


    source_query = """

    SELECT


        ROWID
            AS source_id,


        product_type
            AS hardware,


        name
            AS source


    FROM sources

    """


    df_sources = pd.read_sql_query(
        source_query,
        health_conn
    )


    health_conn.close()


    df_secure["source_id"] = pd.to_numeric(
        df_secure["source_id"],
        errors="coerce"
    ).astype("Int64")


    df_sources["source_id"] = pd.to_numeric(
        df_sources["source_id"],
        errors="coerce"
    ).astype("Int64")


    df_result = df_secure.merge(

        df_sources,

        how="left",

        on="source_id",

        validate="many_to_one"

    )


    return df_result


print(
    "✅ Resting Heart Rate extractor is ready."
)


resting_hr_results = {}


for dataset in target_datasets:


    version = dataset["version"]


    secure_path = dataset[
        "healthdb_secure"
    ]


    healthdb_path = dataset[
        "healthdb"
    ]


    print()

    print("=" * 70)


    print(
        f"❤️‍🩹 RESTING HEART RATE - {version}"
    )


    print("=" * 70)


    try:


        df_resting = extract_resting_heart_rate(

            secure_path,

            healthdb_path

        )


        resting_hr_results[
            version
        ] = df_resting


        print(
            "✅ Extraction successful"
        )


        print(
            f"📊 Total Events: "
            f"{len(df_resting)}"
        )


        source_found = (

            df_resting["source"]
            .notna()
            .sum()

        )


        print(
            f"📱 Events with Source: "
            f"{source_found}"
        )


        print(
            "🔐 Data:"
            " healthdb_secure.sqlite"
        )


        print(
            "📂 Source:"
            " healthdb.sqlite"
        )


    except Exception as e:


        print(
            f"❌ Error: {e}"
        )


for version, df_resting in resting_hr_results.items():


    print()

    print("=" * 70)


    print(
        f"❤️‍🩹 RAW RESTING HEART RATE - {version}"
    )


    print("=" * 70)


    display(
        df_resting.head(20)
    )


    print()


    print(
        f"📊 Total Events: "
        f"{len(df_resting)}"
    )


resting_hr_reports = {}


for version, df_resting in resting_hr_results.items():


    df_report = pd.DataFrame()


    df_report[
        "Start Timestamp"
    ] = (

        df_resting[
            "start_cocoa"
        ]
        .apply(
            parse_cocoa_datetime
        )

    )


    df_report[
        "End Timestamp"
    ] = (

        df_resting[
            "end_cocoa"
        ]
        .apply(
            parse_cocoa_datetime
        )

    )


    df_report[
        "Added Timestamp"
    ] = (

        df_resting[
            "added_cocoa"
        ]
        .apply(
            parse_cocoa_datetime
        )

    )


    df_report[
        "Resting Heart Rate (BPM)"
    ] = (

        df_resting[
            "resting_heart_rate"
        ]

    )


    df_report[
        "Hardware"
    ] = (

        df_resting[
            "hardware"
        ]

    )


    df_report[
        "Source"
    ] = (

        df_resting[
            "source"
        ]

    )


    df_report[
        "Data ID"
    ] = (

        df_resting[
            "data_id"
        ]

    )


    df_report[
        "Health Data Source"
    ] = (
        "healthdb_secure.sqlite"
    )


    df_report[
        "Source Metadata"
    ] = (
        "healthdb.sqlite"
    )


    resting_hr_reports[
        version
    ] = df_report


print(
    "✅ Resting Heart Rate report formatted."
)


for version, df_report in resting_hr_reports.items():


    print()

    print("=" * 70)


    print(
        f"❤️‍🩹 RESTING HEART RATE REPORT - {version}"
    )


    print("=" * 70)


    display(
        df_report.head(20)
    )


    print()


    print(
        f"📊 Total Resting HR Events: "
        f"{len(df_report)}"
    )

# **Source Devices**


print()
print("=" * 70)
print("📱 SOURCE DEVICES VERIFICATION")
print("=" * 70)


for dataset in target_datasets:

    version = dataset["version"]

    healthdb_tables = dataset[
        "healthdb_tables"
    ]


    print()

    print(
        f"📱 {version}"
    )


    if "source_devices" in healthdb_tables:

        print(
            "   📂 source_devices : ✅ Found"
        )

    else:

        print(
            "   📂 source_devices : ❌ Not found"
        )


print()
print("=" * 70)


def get_table_columns(
    database_path,
    table_name
):

    conn = open_sqlite_readonly(
        database_path
    )


    cursor = conn.cursor()


    cursor.execute(
        f"PRAGMA table_info({table_name})"
    )


    columns = [

        row[1]

        for row in cursor.fetchall()

    ]


    conn.close()


    return columns


print(
    "✅ Column inspection is ready."
)


def extract_source_devices(
    health_database_path
):


    columns = get_table_columns(

        health_database_path,

        "source_devices"

    )


    conn = open_sqlite_readonly(
        health_database_path
    )


    if "sync_identity" in columns:


        print(
            "   🔎 Detected schema: iOS 16-17"
        )


        query = """

        SELECT


            creation_date,


            name AS device_name,


            manufacturer,


            model,


            hardware,


            firmware,


            software,


            localIdentifier,


            sync_provenance,


            sync_identity


        FROM source_devices


        WHERE


            name NOT LIKE "__NONE__"


            AND localIdentifier NOT LIKE "__NONE__"


        ORDER BY

            creation_date

        """


    else:


        print(
            "   🔎 Detected schema: iOS 13-15"
        )


        query = """

        SELECT


            creation_date,


            name AS device_name,


            manufacturer,


            model,


            hardware,


            firmware,


            software,


            localIdentifier,


            sync_provenance


        FROM source_devices


        WHERE


            name NOT LIKE "__NONE__"


            AND localIdentifier NOT LIKE "__NONE__"


        ORDER BY

            creation_date

        """


    df = pd.read_sql_query(

        query,

        conn

    )


    conn.close()


    if "sync_identity" not in df.columns:


        df[
            "sync_identity"
        ] = None


    return df


print(
    "✅ Adaptive Source Devices extractor is ready."
)


source_devices_results = {}


for dataset in target_datasets:


    version = dataset["version"]


    healthdb_path = dataset[
        "healthdb"
    ]


    print()

    print("=" * 70)


    print(
        f"📱 SOURCE DEVICES - {version}"
    )


    print("=" * 70)


    try:


        df_devices = extract_source_devices(

            healthdb_path

        )


        source_devices_results[
            version
        ] = df_devices


        print(
            "✅ Extraction successful"
        )


        print(
            f"📊 Total Devices: "
            f"{len(df_devices)}"
        )


        print(
            "📂 Source:"
            " healthdb.sqlite"
        )


    except Exception as e:


        print(
            f"❌ Error: {e}"
        )


for version, df_devices in source_devices_results.items():


    print()

    print("=" * 70)


    print(
        f"📱 RAW SOURCE DEVICES - {version}"
    )


    print("=" * 70)


    display(
        df_devices.head(20)
    )


    print()


    print(
        f"📊 Total Devices: "
        f"{len(df_devices)}"
    )


source_devices_reports = {}


for version, df_devices in source_devices_results.items():


    df_report = pd.DataFrame()


    df_report[
        "Creation Timestamp"
    ] = (

        df_devices[
            "creation_date"
        ]
        .apply(
            parse_cocoa_datetime
        )

    )


    df_report[
        "Device Name"
    ] = (

        df_devices[
            "device_name"
        ]

    )


    df_report[
        "Manufacturer"
    ] = (

        df_devices[
            "manufacturer"
        ]

    )


    df_report[
        "Model"
    ] = (

        df_devices[
            "model"
        ]

    )


    df_report[
        "Hardware"
    ] = (

        df_devices[
            "hardware"
        ]

    )


    df_report[
        "Firmware"
    ] = (

        df_devices[
            "firmware"
        ]

    )


    df_report[
        "Software"
    ] = (

        df_devices[
            "software"
        ]

    )


    df_report[
        "Local Identifier"
    ] = (

        df_devices[
            "localIdentifier"
        ]

    )


    df_report[
        "Sync Provenance"
    ] = (

        df_devices[
            "sync_provenance"
        ]

    )


    df_report[
        "Sync Identity"
    ] = (

        df_devices[
            "sync_identity"
        ]

    )


    df_report[
        "Source Database"
    ] = (

        "healthdb.sqlite"

    )


    source_devices_reports[
        version
    ] = df_report


print(
    "✅ Source Devices report formatted."
)


for version, df_report in source_devices_reports.items():


    print()

    print("=" * 70)


    print(
        f"📱 SOURCE DEVICES REPORT - {version}"
    )


    print("=" * 70)


    display(
        df_report.head(20)
    )


    print()


    print(
        f"📊 Total Devices: "
        f"{len(df_report)}"
    )

# **Steps**


print()
print("=" * 70)
print("🚶 STEPS STRUCTURE VERIFICATION")
print("=" * 70)


steps_required_tables = [
    "samples",
    "quantity_samples",
    "objects",
    "data_provenances"
]


for dataset in target_datasets:

    version = dataset["version"]

    secure_tables = dataset[
        "healthdb_secure_tables"
    ]

    print()
    print(
        f"📱 {version}"
    )

    for table_name in steps_required_tables:

        if table_name in secure_tables:

            print(
                f"   📂 {table_name:<20}: ✅ Found"
            )

        else:

            print(
                f"   📂 {table_name:<20}: ❌ Not found"
            )


print()
print("=" * 70)


def extract_steps(
    secure_database_path
):


    conn = open_sqlite_readonly(
        secure_database_path
    )


    query = """

    SELECT

        samples.start_date AS start_date,

        samples.end_date AS end_date,

        quantity_samples.quantity AS steps,

        data_provenances.origin_product_type AS device

    FROM samples

    LEFT JOIN quantity_samples
        ON samples.data_id =
           quantity_samples.data_id

    LEFT JOIN objects
        ON samples.data_id =
           objects.data_id

    LEFT JOIN data_provenances
        ON objects.provenance =
           data_provenances.rowid

    WHERE

        samples.data_type = 7

        AND quantity_samples.quantity IS NOT NULL

    ORDER BY

        samples.start_date DESC

    """


    df = pd.read_sql_query(
        query,
        conn
    )


    conn.close()


    return df


print(
    "✅ Steps extractor is ready."
)


steps_results = {}


for dataset in target_datasets:

    version = dataset["version"]

    secure_db_path = dataset[
        "healthdb_secure"
    ]


    print()
    print("=" * 70)

    print(
        f"🚶 STEPS - {version}"
    )

    print("=" * 70)


    try:

        df_steps = extract_steps(
            secure_db_path
        )


        steps_results[
            version
        ] = df_steps


        print(
            "✅ Extraction successful"
        )

        print(
            f"📊 Total Step Records: "
            f"{len(df_steps)}"
        )

        print(
            "📂 Source: healthdb_secure.sqlite"
        )


    except Exception as e:

        print(
            f"❌ Error: {e}"
        )


for version, df_steps in steps_results.items():

    print()
    print("=" * 70)

    print(
        f"🚶 RAW STEPS - {version}"
    )

    print("=" * 70)


    display(
        df_steps.head(20)
    )


    print()

    print(
        f"📊 Total Step Records: "
        f"{len(df_steps)}"
    )


steps_reports = {}


for version, df_steps in steps_results.items():

    df_report = pd.DataFrame()


    df_report[
        "Start Timestamp"
    ] = (

        df_steps[
            "start_date"
        ]
        .apply(
            parse_cocoa_datetime
        )

    )


    df_report[
        "End Timestamp"
    ] = (

        df_steps[
            "end_date"
        ]
        .apply(
            parse_cocoa_datetime
        )

    )


    df_report[
        "Duration (Seconds)"
    ] = (

        pd.to_numeric(
            df_steps["end_date"],
            errors="coerce"
        )

        -

        pd.to_numeric(
            df_steps["start_date"],
            errors="coerce"
        )

    )


    df_report[
        "Steps"
    ] = (

        pd.to_numeric(
            df_steps["steps"],
            errors="coerce"
        )

    )


    df_report[
        "Device"
    ] = (

        df_steps[
            "device"
        ]

    )


    df_report[
        "Source Database"
    ] = (

        "healthdb_secure.sqlite"

    )


    steps_reports[
        version
    ] = df_report


print(
    "✅ Steps report formatted."
)


for version, df_report in steps_reports.items():

    print()
    print("=" * 70)

    print(
        f"🚶 STEPS REPORT - {version}"
    )

    print("=" * 70)


    display(
        df_report.head(20)
    )


    print()

    print(
        f"📊 Total Step Records: "
        f"{len(df_report)}"
    )

# **Watch by Sleep Period**


print()
print("=" * 70)
print("⌚ WATCH BY SLEEP PERIOD VERIFICATION")
print("=" * 70)


watch_sleep_required_tables = [
    "samples",
    "category_samples",
    "objects",
    "data_provenances"
]


for dataset in target_datasets:

    version = dataset["version"]

    secure_tables = dataset[
        "healthdb_secure_tables"
    ]

    print()
    print(
        f"📱 {version}"
    )

    for table_name in watch_sleep_required_tables:

        if table_name in secure_tables:

            print(
                f"   📂 {table_name:<20}: ✅ Found"
            )

        else:

            print(
                f"   📂 {table_name:<20}: ❌ Not found"
            )


print()
print("=" * 70)


import re


def get_ios_major_version(version):


    version_text = str(
        version
    )

    match = re.search(
        r"(\d+)",
        version_text
    )

    if not match:

        return None

    return int(
        match.group(1)
    )


print(
    "✅ iOS version detection is ready."
)


def extract_watch_by_sleep_period(
    secure_database_path,
    ios_version
):


    major_version = get_ios_major_version(
        ios_version
    )


    conn = open_sqlite_readonly(
        secure_database_path
    )


    if major_version == 17:

        print(
            "   🔎 Detected sleep schema: "
            "iOS 17 Sleep Stages"
        )


        query = """

        WITH lagged_samples AS (

            SELECT

                s.start_date,

                s.end_date,

                s.data_id,

                s.data_type,

                cs.value,

                LAG(s.data_id)
                    OVER (
                        ORDER BY s.data_id
                    ) AS prev_data_id,

                CASE

                    WHEN cs.value = 2
                        THEN 'AWAKE'

                    WHEN cs.value = 3
                        THEN 'CORE'

                    WHEN cs.value = 4
                        THEN 'DEEP'

                    WHEN cs.value = 5
                        THEN 'REM'

                END AS sleep_value

            FROM samples s

            LEFT JOIN category_samples cs
                ON s.data_id = cs.data_id

            LEFT JOIN objects o
                ON s.data_id = o.data_id

            LEFT JOIN data_provenances dp
                ON o.provenance = dp.rowid

            WHERE

                s.data_type = 63

                AND cs.value NOT IN (0, 1)

                AND dp.origin_product_type
                    LIKE '%Watch%'

        ),


        grouped_samples AS (

            SELECT

                *,

                CASE

                    WHEN data_id - prev_data_id > 1
                        OR prev_data_id IS NULL

                    THEN 1

                    ELSE 0

                END AS is_new_group,


                SUM(

                    CASE

                        WHEN data_id - prev_data_id > 1
                            OR prev_data_id IS NULL

                        THEN 1

                        ELSE 0

                    END

                ) OVER (
                    ORDER BY data_id
                ) AS group_number

            FROM lagged_samples

        )


        SELECT

            MIN(start_date)
                AS start_date,

            MAX(end_date)
                AS end_date,


            SUM(

                CASE

                    WHEN sleep_value IN (
                        'AWAKE',
                        'REM',
                        'CORE',
                        'DEEP'
                    )

                    THEN
                        (end_date - start_date) / 60.0

                    ELSE 0

                END

            ) AS total_duration,


            SUM(

                CASE

                    WHEN sleep_value = 'AWAKE'

                    THEN
                        (end_date - start_date) / 60.0

                    ELSE 0

                END

            ) AS awake_duration,


            SUM(

                CASE

                    WHEN sleep_value = 'REM'

                    THEN
                        (end_date - start_date) / 60.0

                    ELSE 0

                END

            ) AS rem_duration,


            SUM(

                CASE

                    WHEN sleep_value = 'CORE'

                    THEN
                        (end_date - start_date) / 60.0

                    ELSE 0

                END

            ) AS core_duration,


            SUM(

                CASE

                    WHEN sleep_value = 'DEEP'

                    THEN
                        (end_date - start_date) / 60.0

                    ELSE 0

                END

            ) AS deep_duration


        FROM grouped_samples

        GROUP BY group_number

        ORDER BY
            MIN(start_date) ASC

        """


        sleep_structure = (
            "Sleep Stages"
        )


    else:

        print(
            "   🔎 Detected sleep schema: "
            "Legacy IN BED / ASLEEP"
        )


        query = """

        WITH lagged_samples AS (

            SELECT

                s.data_id,

                s.start_date,

                s.end_date,

                (
                    s.end_date -
                    s.start_date
                ) / 60 AS duration_minutes,

                s.data_type,

                cs.value,

                dp.origin_product_type,


                LAG(s.data_id)
                    OVER (
                        ORDER BY s.data_id
                    ) AS prev_data_id,


                CASE cs.value

                    WHEN 0
                        THEN 'IN BED'

                    WHEN 1
                        THEN 'ASLEEP'

                END AS sleep_value


            FROM samples s


            LEFT JOIN category_samples cs

                ON s.data_id =
                   cs.data_id


            LEFT JOIN objects o

                ON s.data_id =
                   o.data_id


            LEFT JOIN data_provenances dp

                ON o.provenance =
                   dp.rowid


            WHERE

                s.data_type = 63

                AND cs.value IN (0, 1)

                AND dp.origin_product_type
                    LIKE '%Watch%'

        ),


        grouped_samples AS (

            SELECT

                *,

                CASE

                    WHEN prev_data_id IS NULL
                        OR data_id - prev_data_id > 1

                    THEN 1

                    ELSE 0

                END AS is_new_group,


                SUM(

                    CASE

                        WHEN prev_data_id IS NULL
                            OR data_id - prev_data_id > 1

                        THEN 1

                        ELSE 0

                    END

                ) OVER (
                    ORDER BY data_id
                ) AS group_number


            FROM lagged_samples

        )


        SELECT

            MIN(start_date)
                AS start_date,

            MAX(end_date)
                AS end_date,

            MAX(data_type)
                AS data_type_id,

            MAX(origin_product_type)
                AS device_name,


            SUM(

                CASE

                    WHEN sleep_value IN (
                        'IN BED',
                        'ASLEEP'
                    )

                    THEN duration_minutes

                    ELSE 0

                END

            ) AS total_duration,


            SUM(

                CASE

                    WHEN sleep_value = 'IN BED'

                    THEN duration_minutes

                    ELSE 0

                END

            ) AS in_bed_duration,

            SUM(

                CASE

                    WHEN sleep_value = 'ASLEEP'

                    THEN duration_minutes

                    ELSE 0

                END

           ) AS asleep_duration,


            CASE

                WHEN SUM(duration_minutes) > 0

                THEN ROUND(

                    SUM(

                        CASE

                            WHEN sleep_value = 'IN BED'

                            THEN duration_minutes

                            ELSE 0

                        END

                    ) * 100.0
                    /
                    SUM(duration_minutes),

                    2

                )

                ELSE 0

            END AS in_bed_percent,


            CASE

                WHEN SUM(duration_minutes) > 0

                THEN ROUND(

                    SUM(

                        CASE

                            WHEN sleep_value = 'ASLEEP'

                            THEN duration_minutes

                            ELSE 0

                        END

                    ) * 100.0
                    /
                    SUM(duration_minutes),

                    2

                )

                ELSE 0

            END AS asleep_percent


        FROM grouped_samples


        GROUP BY
            group_number


        ORDER BY
            MIN(start_date) ASC

        """


        sleep_structure = (
            "In Bed / Asleep"
        )


    df = pd.read_sql_query(
        query,
        conn
    )


    conn.close()


    df[
        "sleep_structure"
    ] = sleep_structure


    return df


print(
    "✅ Adaptive Watch By Sleep Period extractor is ready."
)


watch_sleep_results = {}


for dataset in target_datasets:

    version = dataset[
        "version"
    ]

    secure_db_path = dataset[
        "healthdb_secure"
    ]


    print()
    print("=" * 70)

    print(
        f"⌚ WATCH BY SLEEP PERIOD - {version}"
    )

    print("=" * 70)


    try:

        df_watch_sleep = (
            extract_watch_by_sleep_period(
                secure_db_path,
                version
            )
        )


        watch_sleep_results[
            version
        ] = df_watch_sleep


        print(
            "✅ Extraction successful"
        )

        print(
            f"📊 Total Sleep Periods: "
            f"{len(df_watch_sleep)}"
        )

        print(
            "📂 Source: "
            "healthdb_secure.sqlite"
        )


    except Exception as e:

        print(
            f"❌ Error: {e}"
        )


for version, df_watch_sleep in watch_sleep_results.items():

    print()
    print("=" * 70)

    print(
        f"⌚ RAW WATCH BY SLEEP PERIOD - {version}"
    )

    print("=" * 70)


    display(
        df_watch_sleep.head(20)
    )


    print()

    print(
        f"📊 Total Sleep Periods: "
        f"{len(df_watch_sleep)}"
    )


def sleep_minutes_to_hms(
    minutes_value
):


    try:

        if pd.isna(
            minutes_value
        ):

            return "00:00:00"


        total_seconds = int(
            float(minutes_value) * 60
        )


    except (
        TypeError,
        ValueError
    ):

        return "00:00:00"


    hours = (
        total_seconds // 3600
    )

    minutes = (
        total_seconds % 3600
    ) // 60

    seconds = (
        total_seconds % 60
    )


    return (
        f"{hours:02d}:"
        f"{minutes:02d}:"
        f"{seconds:02d}"
    )


def calculate_sleep_percent(
    part,
    total
):


    try:

        part_value = float(
            part
        )

        total_value = float(
            total
        )


        if total_value == 0:

            return 0.0


        return round(

            (
                part_value
                /
                total_value
            )
            * 100,

            2

        )


    except (
        TypeError,
        ValueError
    ):

        return 0.0


print(
    "✅ Sleep duration and percentage helpers are ready."
)


watch_sleep_reports = {}


for version, df_sleep in watch_sleep_results.items():

    df_report = pd.DataFrame()


    major_version = get_ios_major_version(
        version
    )


    df_report[
        "Start Timestamp"
    ] = (

        df_sleep[
            "start_date"
        ]
        .apply(
            parse_cocoa_datetime
        )

    )


    df_report[
        "End Timestamp"
    ] = (

        df_sleep[
            "end_date"
        ]
        .apply(
            parse_cocoa_datetime
        )

    )


    df_report[
        "Sleep Structure"
    ] = (

        df_sleep[
            "sleep_structure"
        ]

    )


    df_report[
        "Time in Bed (Minutes)"
    ] = (

        pd.to_numeric(
            df_sleep[
                "total_duration"
            ],
            errors="coerce"
        )

    )


    df_report[
        "Time in Bed (HH:MM:SS)"
    ] = (

        df_sleep[
            "total_duration"
        ]
        .apply(
            sleep_minutes_to_hms
        )

    )


    if major_version == 17:


        rem_minutes = pd.to_numeric(
            df_sleep[
                "rem_duration"
            ],
            errors="coerce"
        ).fillna(0)


        core_minutes = pd.to_numeric(
            df_sleep[
                "core_duration"
            ],
            errors="coerce"
        ).fillna(0)


        deep_minutes = pd.to_numeric(
            df_sleep[
                "deep_duration"
            ],
            errors="coerce"
        ).fillna(0)


        time_asleep_minutes = (

            rem_minutes
            +
            core_minutes
            +
            deep_minutes

        )


        df_report[
            "Time Asleep (Minutes)"
        ] = time_asleep_minutes


        df_report[
            "Time Asleep (HH:MM:SS)"
        ] = (

            time_asleep_minutes
            .apply(
                sleep_minutes_to_hms
            )

        )


        total = pd.to_numeric(
            df_sleep[
                "total_duration"
            ],
            errors="coerce"
        )


        df_report[
            "Awake Duration (Minutes)"
        ] = pd.to_numeric(
            df_sleep[
                "awake_duration"
            ],
            errors="coerce"
        )


        df_report[
            "Awake Duration (HH:MM:SS)"
        ] = (

            df_sleep[
                "awake_duration"
            ]
            .apply(
                sleep_minutes_to_hms
            )

        )


        df_report[
            "Awake Percent"
        ] = [

            calculate_sleep_percent(
                part,
                total_value
            )

            for part, total_value in zip(

                df_sleep[
                    "awake_duration"
                ],

                total

            )

        ]


        df_report[
            "REM Duration (Minutes)"
        ] = rem_minutes


        df_report[
            "REM Duration (HH:MM:SS)"
        ] = (

            rem_minutes
            .apply(
                sleep_minutes_to_hms
            )

        )


        df_report[
            "REM Percent"
        ] = [

            calculate_sleep_percent(
                part,
                total_value
            )

            for part, total_value in zip(

                rem_minutes,

                total

            )

        ]


        df_report[
            "Core Duration (Minutes)"
        ] = core_minutes


        df_report[
            "Core Duration (HH:MM:SS)"
        ] = (

            core_minutes
            .apply(
                sleep_minutes_to_hms
            )

        )


        df_report[
            "Core Percent"
        ] = [

            calculate_sleep_percent(
                part,
                total_value
            )

            for part, total_value in zip(

                core_minutes,

                total

            )

        ]


        df_report[
            "Deep Duration (Minutes)"
        ] = deep_minutes


        df_report[
            "Deep Duration (HH:MM:SS)"
        ] = (

            deep_minutes
            .apply(
                sleep_minutes_to_hms
            )

        )


        df_report[
            "Deep Percent"
        ] = [

            calculate_sleep_percent(
                part,
                total_value
            )

            for part, total_value in zip(

                deep_minutes,

                total

            )

        ]


        df_report[
            "Data Type ID"
        ] = None


        df_report[
            "Device"
        ] = None


        df_report[
            "In Bed Duration (Minutes)"
        ] = None


        df_report[
            "In Bed Duration (HH:MM:SS)"
        ] = None


        df_report[
            "In Bed Percent"
        ] = None


        df_report[
            "Asleep Percent"
        ] = (

            [

                calculate_sleep_percent(
                    asleep,
                    total_value
                )

                for asleep, total_value in zip(

                    time_asleep_minutes,

                    total

                )

            ]

        )


    else:


        asleep_minutes = pd.to_numeric(
            df_sleep[
                "asleep_duration"
            ],
            errors="coerce"
        )


        df_report[
            "Time Asleep (Minutes)"
        ] = asleep_minutes


        df_report[
            "Time Asleep (HH:MM:SS)"
        ] = (

            asleep_minutes
            .apply(
                sleep_minutes_to_hms
            )

        )


        df_report[
            "Data Type ID"
        ] = (

            df_sleep[
                "data_type_id"
            ]

        )


        df_report[
            "Device"
        ] = (

            df_sleep[
                "device_name"
            ]

        )


        df_report[
            "In Bed Duration (Minutes)"
        ] = (

            pd.to_numeric(
                df_sleep[
                    "in_bed_duration"
                ],
                errors="coerce"
            )

        )


        df_report[
            "In Bed Duration (HH:MM:SS)"
        ] = (

            df_sleep[
                "in_bed_duration"
            ]
            .apply(
                sleep_minutes_to_hms
            )

        )


        df_report[
            "In Bed Percent"
        ] = (

            df_sleep[
                "in_bed_percent"
            ]

        )


        df_report[
            "Asleep Percent"
        ] = (

            df_sleep[
                "asleep_percent"
            ]

        )


        df_report[
            "Awake Duration (Minutes)"
        ] = None


        df_report[
            "Awake Duration (HH:MM:SS)"
        ] = None


        df_report[
            "Awake Percent"
        ] = None


        df_report[
            "REM Duration (Minutes)"
        ] = None


        df_report[
            "REM Duration (HH:MM:SS)"
        ] = None


        df_report[
            "REM Percent"
        ] = None


        df_report[
            "Core Duration (Minutes)"
        ] = None


        df_report[
            "Core Duration (HH:MM:SS)"
        ] = None


        df_report[
            "Core Percent"
        ] = None


        df_report[
            "Deep Duration (Minutes)"
        ] = None


        df_report[
            "Deep Duration (HH:MM:SS)"
        ] = None


        df_report[
            "Deep Percent"
        ] = None


    df_report[
        "Source Database"
    ] = (
        "healthdb_secure.sqlite"
    )


    watch_sleep_reports[
        version
    ] = df_report


print(
    "✅ Watch By Sleep Period report formatted."
)


for version, df_report in watch_sleep_reports.items():

    print()
    print("=" * 70)

    print(
        f"⌚ WATCH BY SLEEP PERIOD REPORT - {version}"
    )

    print("=" * 70)


    display(
        df_report.head(20)
    )


    print()

    print(
        f"📊 Total Sleep Periods: "
        f"{len(df_report)}"
    )

# **Watch Worn Data**


print()
print("=" * 70)
print("⌚ WATCH WORN DATA VERIFICATION")
print("=" * 70)


for dataset in target_datasets:

    version = dataset[
        "version"
    ]

    secure_tables = dataset[
        "healthdb_secure_tables"
    ]


    print()
    print(
        f"📱 {version}"
    )


    if "samples" in secure_tables:

        print(
            "   📂 samples              : ✅ Found"
        )

    else:

        print(
            "   📂 samples              : ❌ Not found"
        )


print()
print("=" * 70)


def extract_watch_worn(
    secure_database_path
):


    conn = open_sqlite_readonly(
        secure_database_path
    )


    query = """

    WITH TimeData AS (

        SELECT

            s.start_date AS start_cocoa,

            s.end_date AS end_cocoa,

            LAG(
                s.end_date
            ) OVER (
                ORDER BY s.start_date
            ) AS prev_end_cocoa

        FROM samples s

        WHERE
            s.data_type IN (70, '70')

    ),


    PeriodData AS (

        SELECT

            *,

            (
                start_cocoa
                -
                prev_end_cocoa
            ) AS gap_seconds,


            CASE

                WHEN (
                    start_cocoa
                    -
                    prev_end_cocoa
                ) > 3600

                THEN 1

                ELSE 0

            END AS new_period

        FROM TimeData

    ),


    PeriodGroup AS (

        SELECT

            *,

            SUM(
                new_period
            ) OVER (

                ORDER BY start_cocoa

                ROWS BETWEEN
                    UNBOUNDED PRECEDING
                    AND CURRENT ROW

            ) AS period_id

        FROM PeriodData

    ),


    Summary AS (

        SELECT

            period_id,


            MIN(
                start_cocoa
            ) AS period_start_cocoa,


            MAX(
                end_cocoa
            ) AS period_end_cocoa,


            CAST(

                (
                    MAX(end_cocoa)
                    -
                    MIN(start_cocoa)
                ) / 3600

                AS INT

            ) AS hours_worn


        FROM PeriodGroup


        GROUP BY
            period_id

    )


    SELECT

        s1.period_start_cocoa
            AS start_cocoa,


        s1.hours_worn
            AS hours_worn,


        s1.period_end_cocoa
            AS end_cocoa,


        CAST(

            (
                s2.period_start_cocoa
                -
                s1.period_end_cocoa
            ) / 3600

            AS INT

        ) AS hours_off_before_next


    FROM Summary s1


    LEFT JOIN Summary s2

        ON s1.period_id + 1 =
           s2.period_id


    ORDER BY
        s1.period_id

    """


    df = pd.read_sql_query(
        query,
        conn
    )


    conn.close()


    return df


print(
    "✅ Watch Worn Data extractor is ready."
)


watch_worn_results = {}


for dataset in target_datasets:

    version = dataset[
        "version"
    ]

    secure_db_path = dataset[
        "healthdb_secure"
    ]


    print()
    print("=" * 70)

    print(
        f"⌚ WATCH WORN DATA - {version}"
    )

    print("=" * 70)


    try:

        df_watch_worn = extract_watch_worn(
            secure_db_path
        )


        watch_worn_results[
            version
        ] = df_watch_worn


        print(
            "✅ Extraction successful"
        )

        print(
            f"📊 Total Worn Periods: "
            f"{len(df_watch_worn)}"
        )

        print(
            "📂 Source: "
            "healthdb_secure.sqlite"
        )


    except Exception as e:

        print(
            f"❌ Error: {e}"
        )


for version, df_watch_worn in watch_worn_results.items():

    print()
    print("=" * 70)

    print(
        f"⌚ RAW WATCH WORN DATA - {version}"
    )

    print("=" * 70)


    display(
        df_watch_worn.head(20)
    )


    print()

    print(
        f"📊 Total Worn Periods: "
        f"{len(df_watch_worn)}"
    )


watch_worn_reports = {}


for version, df_watch_worn in watch_worn_results.items():

    df_report = pd.DataFrame()


    df_report[
        "Start Time"
    ] = (

        df_watch_worn[
            "start_cocoa"
        ]
        .apply(
            parse_cocoa_datetime
        )

    )


    df_report[
        "Last Worn Time"
    ] = (

        df_watch_worn[
            "end_cocoa"
        ]
        .apply(
            parse_cocoa_datetime
        )

    )


    df_report[
        "Hours Worn"
    ] = (

        pd.to_numeric(
            df_watch_worn[
                "hours_worn"
            ],
            errors="coerce"
        )

    )


    df_report[
        "Hours Off Before Next"
    ] = (

        pd.to_numeric(
            df_watch_worn[
                "hours_off_before_next"
            ],
            errors="coerce"
        )

    )


    df_report[
        "Source Database"
    ] = (
        "healthdb_secure.sqlite"
    )


    watch_worn_reports[
        version
    ] = df_report


print(
    "✅ Watch Worn Data report formatted."
)


for version, df_report in watch_worn_reports.items():

    print()
    print("=" * 70)

    print(
        f"⌚ WATCH WORN DATA REPORT - {version}"
    )

    print("=" * 70)


    display(
        df_report.head(20)
    )


    print()

    print(
        f"📊 Total Worn Periods: "
        f"{len(df_report)}"
    )

# **Weight**


print()
print("=" * 70)
print("⚖️ WEIGHT STRUCTURE VERIFICATION")
print("=" * 70)


weight_required_tables = [
    "samples",
    "quantity_samples"
]


for dataset in target_datasets:

    version = dataset[
        "version"
    ]

    secure_tables = dataset[
        "healthdb_secure_tables"
    ]


    print()
    print(
        f"📱 {version}"
    )


    for table_name in weight_required_tables:

        if table_name in secure_tables:

            print(
                f"   📂 {table_name:<20}: ✅ Found"
            )

        else:

            print(
                f"   📂 {table_name:<20}: ❌ Not found"
            )


print()
print("=" * 70)


def extract_weight(
    secure_database_path
):


    conn = open_sqlite_readonly(
        secure_database_path
    )


    query = """

    SELECT

        t.start_date
            AS start_date,

        t.kg_str
            AS weight_kg_str,

        t.stone_int
            AS stone_int,

        t.pounds_int
            AS pounds_int,

        t.lbs_str
            AS weight_lbs_str


    FROM (

        SELECT

            s.start_date
                AS start_date,


            substr(
                qs.quantity,
                1,
                5
            ) AS kg_str,


            CAST(
                qs.quantity / 6.35029317
                AS INT
            ) AS stone_int,


            CAST(

                (
                    (
                        (
                            qs.quantity / 6.35029317
                        )
                        -
                        CAST(
                            qs.quantity / 6.35029317
                            AS INT
                        )
                    )
                    * 14
                )
                + 0.5

                AS INT

            ) AS pounds_int,


            substr(
                qs.quantity * 2.20462262,
                1,
                6
            ) AS lbs_str


        FROM samples s


        JOIN quantity_samples qs

            ON s.data_id =
               qs.data_id


        WHERE

            s.data_type IN (3, '3')

            AND qs.quantity IS NOT NULL

    ) t


    ORDER BY
        t.start_date DESC

    """


    df = pd.read_sql_query(
        query,
        conn
    )


    conn.close()


    return df


print(
    "✅ Weight extractor is ready."
)


weight_results = {}


for dataset in target_datasets:

    version = dataset[
        "version"
    ]

    secure_db_path = dataset[
        "healthdb_secure"
    ]


    print()
    print("=" * 70)

    print(
        f"⚖️ WEIGHT - {version}"
    )

    print("=" * 70)


    try:

        df_weight = extract_weight(
            secure_db_path
        )


        weight_results[
            version
        ] = df_weight


        print(
            "✅ Extraction successful"
        )

        print(
            f"📊 Total Weight Records: "
            f"{len(df_weight)}"
        )

        print(
            "📂 Source: "
            "healthdb_secure.sqlite"
        )


    except Exception as e:

        print(
            f"❌ Error: {e}"
        )


for version, df_weight in weight_results.items():

    print()
    print("=" * 70)

    print(
        f"⚖️ RAW WEIGHT - {version}"
    )

    print("=" * 70)


    display(
        df_weight.head(20)
    )


    print()

    print(
        f"📊 Total Weight Records: "
        f"{len(df_weight)}"
    )


def format_stone_pounds(
    stone_value,
    pounds_value
):


    try:

        stone = (
            int(stone_value)
            if pd.notna(stone_value)
            else 0
        )

        pounds = (
            int(pounds_value)
            if pd.notna(pounds_value)
            else 0
        )


        if pounds == 14:

            stone += 1

            pounds = 0


        return (
            f"{stone} Stone "
            f"{pounds} Pounds"
        )


    except (
        TypeError,
        ValueError
    ):

        return None


print(
    "✅ Stone/Pounds helper is ready."
)


weight_reports = {}


for version, df_weight in weight_results.items():

    df_report = pd.DataFrame()


    df_report[
        "Timestamp"
    ] = (

        df_weight[
            "start_date"
        ]
        .apply(
            parse_cocoa_datetime
        )

    )


    df_report[
        "Weight (kg)"
    ] = (

        df_weight[
            "weight_kg_str"
        ]

    )


    df_report[
        "Weight (lbs)"
    ] = (

        df_weight[
            "weight_lbs_str"
        ]

    )


    df_report[
        "Weight (Stone/Pounds)"
    ] = [

        format_stone_pounds(
            stone,
            pounds
        )

        for stone, pounds in zip(

            df_weight[
                "stone_int"
            ],

            df_weight[
                "pounds_int"
            ]

        )

    ]


    df_report[
        "Source Database"
    ] = (
        "healthdb_secure.sqlite"
    )


    weight_reports[
        version
    ] = df_report


print(
    "✅ Weight report formatted."
)


for version, df_report in weight_reports.items():

    print()
    print("=" * 70)

    print(
        f"⚖️ WEIGHT REPORT - {version}"
    )

    print("=" * 70)


    display(
        df_report.head(20)
    )


    print()

    print(
        f"📊 Total Weight Records: "
        f"{len(df_report)}"
    )

# **Workouts**


print()
print("=" * 70)
print("🏋️ WORKOUTS STRUCTURE VERIFICATION")
print("=" * 70)


legacy_secure_tables = [
    "workouts",
    "samples"
]


latest_secure_tables = [
    "workout_activities",
    "workouts",
    "workout_statistics",
    "samples",
    "metadata_keys",
    "metadata_values",
    "objects",
    "data_provenances"
]


latest_healthdb_tables = [
    "source_devices",
    "sources"
]


for dataset in target_datasets:

    version = dataset["version"]

    major_version = get_ios_major_version(
        version
    )

    secure_tables = dataset[
        "healthdb_secure_tables"
    ]

    healthdb_tables = dataset[
        "healthdb_tables"
    ]


    print()
    print(
        f"📱 {version}"
    )


    if major_version in (16, 17):

        print(
            "   🔎 Schema: "
            "Workout Activities iOS 16-17"
        )

        print(
            "   📂 healthdb_secure.sqlite"
        )

        for table_name in latest_secure_tables:

            if table_name in secure_tables:

                print(
                    f"      {table_name:<22}: ✅"
                )

            else:

                print(
                    f"      {table_name:<22}: ❌"
                )


        print(
            "   📂 healthdb.sqlite"
        )

        for table_name in latest_healthdb_tables:

            if table_name in healthdb_tables:

                print(
                    f"      {table_name:<22}: ✅"
                )

            else:

                print(
                    f"      {table_name:<22}: ❌"
                )


    else:

        print(
            "   🔎 Schema: "
            "Legacy Workouts iOS 13-15"
        )

        print(
            "   📂 healthdb_secure.sqlite"
        )

        for table_name in legacy_secure_tables:

            if table_name in secure_tables:

                print(
                    f"      {table_name:<22}: ✅"
                )

            else:

                print(
                    f"      {table_name:<22}: ❌"
                )


print()
print("=" * 70)


WORKOUT_ACTIVITY_LABELS = {

    1: "AMERICAN FOOTBALL",
    2: "ARCHERY",
    3: "AUSTRALIAN FOOTBALL",
    4: "BADMINTON",
    5: "BASEBALL",
    6: "BASKETBALL",
    7: "BOWLING",
    8: "BOXING",
    9: "CLIMBING",
    10: "CRICKET",
    11: "CROSS TRAINING",
    12: "CURLING",
    13: "CYCLING",

    16: "ELLIPTICAL",
    17: "EQUESTRIAN SPORTS",
    18: "FENCING",
    19: "FISHING",
    20: "FUNCTION STRENGTH TRAINING",
    21: "GOLF",
    22: "GYMNASTICS",
    23: "HANDBALL",
    24: "HIKING",
    25: "HOCKEY",
    26: "HUNTING",
    27: "LACROSS",
    28: "MARTIAL ARTS",
    29: "MIND AND BODY",

    31: "PADDLE SPORTS",
    32: "PLAY",
    33: "PREPARATION AND RECOVERY",
    34: "RACQUETBALL",
    35: "ROWING",
    36: "RUGBY",
    37: "RUNNING",
    38: "SAILING",
    39: "SKATING SPORTS",
    40: "SNOW SPORTS",
    41: "SOCCER",
    42: "SOFTBALL",
    43: "SQUASH",
    44: "STAIRSTEPPER",
    45: "SURFING SPORTS",
    46: "SWIMMING",
    47: "TABLE TENNIS",
    48: "TENNIS",
    49: "TRACK AND FIELD",
    50: "TRADITIONAL STRENGTH TRAINING",
    51: "VOLLEYBALL",
    52: "WALKING",
    53: "WATER FITNESS",
    54: "WATER POLO",
    55: "WATER SPORTS",
    56: "WRESTLING",
    57: "YOGA",
    58: "BARRE",
    59: "CORE TRAINING",
    60: "CROSS COUNTRY SKIING",
    61: "DOWNHILL SKIING",
    62: "FLEXIBILITY",
    63: "HIGH INTENSITY INTERVAL TRAINING (HIIT)",
    64: "JUMP ROPE",
    65: "KICKBOXING",
    66: "PILATES",
    67: "SNOWBOARDING",
    68: "STAIRS",
    69: "STEP TRAINING",
    70: "WHEELCHAIR WALK PACE",
    71: "WHEELCHAIR RUN PACE",
    72: "TAI CHI",
    73: "MIXED CARDIO",
    74: "HAND CYCLING",
    75: "DISC SPORTS",
    76: "FITNESS GAMING",
    77: "DANCE",
    78: "SOCIAL DANCE",
    79: "PICKLEBALL",
    80: "COOLDOWN",

    3000: "OTHER"
}


WORKOUT_GOAL_LABELS = {

    0: "Open",
    1: "Distance in meters",
    2: "Time in seconds",
    3: "Kilocalories"

}


WORKOUT_LOCATION_LABELS = {

    2: "Indoor",
    3: "Outdoor"

}


def get_workout_enum_label(
    value,
    mapping,
    default_value
):


    if value is None:

        return default_value


    try:

        if pd.isna(value):

            return default_value

    except TypeError:

        pass


    try:

        key = int(
            float(value)
        )

    except (
        TypeError,
        ValueError
    ):

        return default_value


    return mapping.get(
        key,
        default_value
    )


print(
    "✅ Workout enumeration mapping is ready."
)


def open_workouts_cross_database_readonly(
    secure_database_path,
    health_database_path=None
):


    conn = sqlite3.connect(
        ":memory:",
        uri=True
    )


    secure_uri = (
        Path(
            secure_database_path
        )
        .resolve()
        .as_uri()
        +
        "?mode=ro"
    )


    conn.execute(
        "ATTACH DATABASE ? AS secure_db",
        (
            secure_uri,
        )
    )


    if health_database_path is not None:

        health_uri = (
            Path(
                health_database_path
            )
            .resolve()
            .as_uri()
            +
            "?mode=ro"
        )


        conn.execute(
            "ATTACH DATABASE ? AS health_db",
            (
                health_uri,
            )
        )


    return conn


print(
    "✅ Workout cross-database connection is ready."
)


def extract_workouts(
    secure_database_path,
    health_database_path,
    ios_version
):


    major_version = get_ios_major_version(
        ios_version
    )


    if major_version in (16, 17):

        print(
            "   🔎 Detected schema: "
            "iOS 16-17 Workout Activities"
        )


        conn = open_workouts_cross_database_readonly(
            secure_database_path,
            health_database_path
        )


        query = """

        SELECT

            wa.start_date
                AS start_cocoa,

            wa.end_date
                AS end_cocoa,

            wa.activity_type
                AS activity_type_code,

            wa.location_type
                AS location_type_code,

            wa.duration
                AS duration_seconds,


            s.start_date
                AS sample_start_cocoa,

            s.end_date
                AS sample_end_cocoa,


            w.total_distance
                AS total_distance_km,

            w.goal_type
                AS goal_type_code,

            w.goal
                AS goal_value,


            MAX(

                CASE

                    WHEN ws.data_type = 10

                    THEN ROUND(
                        ws.quantity,
                        2
                    )

                END

            ) AS total_active_kcal,


            MAX(

                CASE

                    WHEN ws.data_type = 9

                    THEN ROUND(
                        ws.quantity,
                        2
                    )

                END

            ) AS total_resting_kcal,


            MAX(

                CASE

                    WHEN mk.key = 'HKAverageMETs'

                    THEN ROUND(
                        mv.numerical_value,
                        1
                    )

                END

            ) AS avg_mets,


            MAX(

                CASE

                    WHEN mk.key =
                    '_HKPrivateWorkoutMinHeartRate'

                    THEN CAST(

                        ROUND(
                            mv.numerical_value * 60
                        )

                        AS INT

                    )

                END

            ) AS min_hr_bpm,


            MAX(

                CASE

                    WHEN mk.key =
                    '_HKPrivateWorkoutMaxHeartRate'

                    THEN CAST(

                        ROUND(
                            mv.numerical_value * 60
                        )

                        AS INT

                    )

                END

            ) AS max_hr_bpm,


            MAX(

                CASE

                    WHEN mk.key =
                    '_HKPrivateWorkoutAverageHeartRate'

                    THEN CAST(

                        ROUND(
                            mv.numerical_value * 60
                        )

                        AS INT

                    )

                END

            ) AS avg_hr_bpm,


            MAX(

                CASE

                    WHEN mk.key =
                    'HKWeatherTemperature'

                    THEN ROUND(
                        mv.numerical_value,
                        2
                    )

                END

            ) AS temp_f,


            MAX(

                CASE

                    WHEN mk.key =
                    'HKWeatherHumidity'

                    THEN CAST(
                        mv.numerical_value
                        AS INT
                    )

                END

            ) AS humidity_pct,


            MAX(

                CASE

                    WHEN mk.key =
                    '_HKPrivateWorkoutWeatherLocationCoordinatesLatitude'

                    THEN mv.numerical_value

                END

            ) AS lat,


            MAX(

                CASE

                    WHEN mk.key =
                    '_HKPrivateWorkoutWeatherLocationCoordinatesLongitude'

                    THEN mv.numerical_value

                END

            ) AS lon,


            MAX(

                CASE

                    WHEN mk.key =
                    '_HKPrivateWorkoutMinGroundElevation'

                    THEN ROUND(
                        mv.numerical_value,
                        2
                    )

                END

            ) AS min_elev_m,


            MAX(

                CASE

                    WHEN mk.key =
                    '_HKPrivateWorkoutMaxGroundElevation'

                    THEN ROUND(
                        mv.numerical_value,
                        2
                    )

                END

            ) AS max_elev_m,


            sd.hardware
                AS hardware,


            src.name
                AS source_name,


            dp.source_version
                AS software_version,


            dp.tz_name
                AS timezone_name,


            o.creation_date
                AS added_cocoa


        FROM secure_db.workout_activities AS wa


        LEFT JOIN secure_db.workouts AS w

            ON w.data_id =
               wa.owner_id


        LEFT JOIN secure_db.workout_statistics AS ws

            ON ws.workout_activity_id =
               wa.ROWID


        LEFT JOIN secure_db.samples AS s

            ON s.data_id =
               w.data_id


        LEFT JOIN secure_db.metadata_values AS mv

            ON mv.object_id =
               wa.owner_id


        LEFT JOIN secure_db.metadata_keys AS mk

            ON mk.ROWID =
               mv.key_id


        LEFT JOIN secure_db.objects AS o

            ON o.data_id =
               wa.owner_id


        LEFT JOIN secure_db.data_provenances AS dp

            ON dp.ROWID =
               o.provenance


        LEFT JOIN health_db.source_devices AS sd

            ON sd.ROWID =
               dp.device_id


        LEFT JOIN health_db.sources AS src

            ON src.ROWID =
               dp.source_id


        GROUP BY
            wa.ROWID


        ORDER BY
            wa.start_date

        """


        df = pd.read_sql_query(
            query,
            conn
        )


        df[
            "workout_schema"
        ] = (
            "Workout Activities iOS 16-17"
        )


        conn.close()


        return df


    print(
        "   🔎 Detected schema: "
        "iOS 13-15 Legacy Workouts"
    )


    conn = open_workouts_cross_database_readonly(
        secure_database_path
    )


    query = """

    SELECT

        samples.start_date
            AS start_date,

        samples.end_date
            AS end_date,

        workouts.activity_type
            AS activity_type,

        workouts.duration
            AS duration,

        workouts.total_distance
            AS total_distance,

        workouts.total_energy_burned
            AS total_energy_burned,

        workouts.total_basal_energy_burned
            AS total_basal_energy_burned,

        workouts.goal_type
            AS goal_type,

        workouts.goal
            AS goal

    FROM secure_db.workouts AS workouts


    LEFT JOIN secure_db.samples AS samples

        ON samples.data_id =
           workouts.data_id


    GROUP BY
        workouts.data_id


    ORDER BY
        samples.start_date

    """


    df = pd.read_sql_query(
        query,
        conn
    )


    df[
        "workout_schema"
    ] = (
        "Legacy Workouts iOS 13-15"
    )


    conn.close()


    return df


print(
    "✅ Adaptive Workouts extractor is ready."
)


workouts_results = {}


for dataset in target_datasets:

    version = dataset[
        "version"
    ]


    print()
    print("=" * 70)

    print(
        f"🏋️ WORKOUTS - {version}"
    )

    print("=" * 70)


    try:

        secure_db_path = dataset[
            "healthdb_secure"
        ]


        major_version = get_ios_major_version(
            version
        )


        if major_version in (16, 17):

            healthdb_path = dataset[
                "healthdb"
            ]


        else:

            healthdb_path = None


        df_workouts = extract_workouts(

            secure_db_path,

            healthdb_path,

            version

        )


        workouts_results[
            version
        ] = df_workouts


        print(
            "✅ Extraction successful"
        )

        print(
            f"📊 Total Workouts: "
            f"{len(df_workouts)}"
        )


        if major_version in (16, 17):

            print(
                "📂 Source: "
                "healthdb_secure.sqlite "
                "+ healthdb.sqlite"
            )

        else:

            print(
                "📂 Source: "
                "healthdb_secure.sqlite"
            )


    except Exception as e:

        print(
            f"❌ Error {version}: "
            f"{type(e).__name__}: {e}"
        )


for version, df_workouts in workouts_results.items():

    print()
    print("=" * 70)

    print(
        f"🏋️ RAW WORKOUTS - {version}"
    )

    print("=" * 70)


    display(
        df_workouts.head(20)
    )


    print()

    print(
        f"📊 Total Workouts: "
        f"{len(df_workouts)}"
    )


def workout_seconds_to_hms(
    seconds_value
):


    try:

        if pd.isna(
            seconds_value
        ):

            return None


        total = int(
            float(seconds_value)
        )


    except (
        TypeError,
        ValueError
    ):

        return None


    hours = (
        total // 3600
    )

    minutes = (
        total % 3600
    ) // 60

    seconds = (
        total % 60
    )


    return (
        f"{hours:02d}:"
        f"{minutes:02d}:"
        f"{seconds:02d}"
    )


def workout_km_to_miles(
    km_value
):


    try:

        if pd.isna(
            km_value
        ):

            return None


        return round(
            float(km_value)
            * 0.621371,
            2
        )


    except (
        TypeError,
        ValueError
    ):

        return None


def workout_fahrenheit_to_celsius(
    temp_f
):


    try:

        if pd.isna(
            temp_f
        ):

            return None


        return round(

            (
                float(temp_f)
                -
                32.0
            )
            *
            (
                5.0 / 9.0
            ),

            2

        )


    except (
        TypeError,
        ValueError
    ):

        return None


print(
    "✅ Workout formatting helpers are ready."
)


workouts_reports = {}


for version, df_workouts in workouts_results.items():

    df_report = pd.DataFrame()


    major_version = get_ios_major_version(
        version
    )


    if major_version in (16, 17):


        df_report[
            "Start Timestamp"
        ] = (

            df_workouts[
                "start_cocoa"
            ]
            .apply(
                parse_cocoa_datetime
            )

        )


        df_report[
            "End Timestamp"
        ] = (

            df_workouts[
                "end_cocoa"
            ]
            .apply(
                parse_cocoa_datetime
            )

        )


        df_report[
            "Added Timestamp"
        ] = (

            df_workouts[
                "added_cocoa"
            ]
            .apply(
                parse_cocoa_datetime
            )

        )


        df_report[
            "Activity Type Code"
        ] = (

            df_workouts[
                "activity_type_code"
            ]

        )


        df_report[
            "Activity"
        ] = [

            get_workout_enum_label(
                value,
                WORKOUT_ACTIVITY_LABELS,
                "OTHER"
            )

            for value in df_workouts[
                "activity_type_code"
            ]

        ]


        df_report[
            "Location Type Code"
        ] = (

            df_workouts[
                "location_type_code"
            ]

        )


        df_report[
            "Location"
        ] = [

            get_workout_enum_label(
                value,
                WORKOUT_LOCATION_LABELS,
                "Unknown"
            )

            for value in df_workouts[
                "location_type_code"
            ]

        ]


        duration_seconds = pd.to_numeric(

            df_workouts[
                "duration_seconds"
            ],

            errors="coerce"

        )


        df_report[
            "Workout Duration"
        ] = (

            duration_seconds
            .apply(
                workout_seconds_to_hms
            )

        )


        sample_start = pd.to_numeric(

            df_workouts[
                "sample_start_cocoa"
            ],

            errors="coerce"

        )


        sample_end = pd.to_numeric(

            df_workouts[
                "sample_end_cocoa"
            ],

            errors="coerce"

        )


        sample_duration = (

            sample_end
            -
            sample_start

        )


        df_report[
            "Total Time Duration"
        ] = (

            sample_duration
            .apply(
                workout_seconds_to_hms
            )

        )


        distance_km = pd.to_numeric(

            df_workouts[
                "total_distance_km"
            ],

            errors="coerce"

        )


        df_report[
            "Total Distance (km)"
        ] = (

            distance_km.round(
                2
            )

        )


        df_report[
            "Total Distance (miles)"
        ] = (

            distance_km
            .apply(
                workout_km_to_miles
            )

        )


        df_report[
            "Goal Type Code"
        ] = (

            df_workouts[
                "goal_type_code"
            ]

        )


        df_report[
            "Goal Type"
        ] = [

            get_workout_enum_label(
                value,
                WORKOUT_GOAL_LABELS,
                "Unknown"
            )

            for value in df_workouts[
                "goal_type_code"
            ]

        ]


        df_report[
            "Goal"
        ] = (

            df_workouts[
                "goal_value"
            ]

        )


        df_report[
            "Total Active Energy (kcal)"
        ] = (

            pd.to_numeric(
                df_workouts[
                    "total_active_kcal"
                ],
                errors="coerce"
            )

        )


        df_report[
            "Total Resting Energy (kcal)"
        ] = (

            pd.to_numeric(
                df_workouts[
                    "total_resting_kcal"
                ],
                errors="coerce"
            )

        )


        df_report[
            "Average METs"
        ] = (

            df_workouts[
                "avg_mets"
            ]

        )


        df_report[
            "Min Heart Rate (BPM)"
        ] = (

            df_workouts[
                "min_hr_bpm"
            ]

        )


        df_report[
            "Max Heart Rate (BPM)"
        ] = (

            df_workouts[
                "max_hr_bpm"
            ]

        )


        df_report[
            "Average Heart Rate (BPM)"
        ] = (

            df_workouts[
                "avg_hr_bpm"
            ]

        )


        df_report[
            "Temperature (F)"
        ] = (

            df_workouts[
                "temp_f"
            ]

        )


        df_report[
            "Temperature (C)"
        ] = (

            df_workouts[
                "temp_f"
            ]
            .apply(
                workout_fahrenheit_to_celsius
            )

        )


        df_report[
            "Humidity (%)"
        ] = (

            df_workouts[
                "humidity_pct"
            ]

        )


        df_report[
            "Latitude"
        ] = (

            df_workouts[
                "lat"
            ]

        )


        df_report[
            "Longitude"
        ] = (

            df_workouts[
                "lon"
            ]

        )


        df_report[
            "Min Ground Elevation (m)"
        ] = (

            df_workouts[
                "min_elev_m"
            ]

        )


        df_report[
            "Max Ground Elevation (m)"
        ] = (

            df_workouts[
                "max_elev_m"
            ]

        )


        df_report[
            "Hardware"
        ] = (

            df_workouts[
                "hardware"
            ]

        )


        df_report[
            "Source"
        ] = (

            df_workouts[
                "source_name"
            ]

        )


        df_report[
            "Software Version"
        ] = (

            df_workouts[
                "software_version"
            ]

        )


        df_report[
            "Timezone"
        ] = (

            df_workouts[
                "timezone_name"
            ]

        )


        df_report[
            "Total Energy Burned (kcal)"
        ] = None

        df_report[
            "Total Basal Energy Burned (kcal)"
        ] = None


    else:


        df_report[
            "Start Timestamp"
        ] = (

            df_workouts[
                "start_date"
            ]
            .apply(
                parse_cocoa_datetime
            )

        )


        df_report[
            "End Timestamp"
        ] = (

            df_workouts[
                "end_date"
            ]
            .apply(
                parse_cocoa_datetime
            )

        )


        df_report[
            "Added Timestamp"
        ] = None


        df_report[
            "Activity Type Code"
        ] = (

            df_workouts[
                "activity_type"
            ]

        )


        df_report[
            "Activity"
        ] = [

            get_workout_enum_label(
                value,
                WORKOUT_ACTIVITY_LABELS,
                "OTHER"
            )

            for value in df_workouts[
                "activity_type"
            ]

        ]


        df_report[
            "Location Type Code"
        ] = None

        df_report[
            "Location"
        ] = None


        duration_seconds = pd.to_numeric(

            df_workouts[
                "duration"
            ],

            errors="coerce"

        )


        df_report[
            "Workout Duration"
        ] = (

            duration_seconds
            .apply(
                workout_seconds_to_hms
            )

        )


        df_report[
            "Total Time Duration"
        ] = None


        distance_km = pd.to_numeric(

            df_workouts[
                "total_distance"
            ],

            errors="coerce"

        )


        df_report[
            "Total Distance (km)"
        ] = (

            distance_km
            .round(2)

        )


        df_report[
            "Total Distance (miles)"
        ] = (

            distance_km
            .apply(
                workout_km_to_miles
            )

        )


        df_report[
            "Goal Type Code"
        ] = (

            df_workouts[
                "goal_type"
            ]

        )


        df_report[
            "Goal Type"
        ] = [

            get_workout_enum_label(
                value,
                WORKOUT_GOAL_LABELS,
                "Unknown"
            )

            for value in df_workouts[
                "goal_type"
            ]

        ]


        df_report[
            "Goal"
        ] = (

            df_workouts[
                "goal"
            ]

        )


        df_report[
            "Total Energy Burned (kcal)"
        ] = (

            df_workouts[
                "total_energy_burned"
            ]

        )


        df_report[
            "Total Basal Energy Burned (kcal)"
        ] = (

            df_workouts[
                "total_basal_energy_burned"
            ]

        )


        df_report[
            "Total Active Energy (kcal)"
        ] = None

        df_report[
            "Total Resting Energy (kcal)"
        ] = None

        df_report[
            "Average METs"
        ] = None

        df_report[
            "Min Heart Rate (BPM)"
        ] = None

        df_report[
            "Max Heart Rate (BPM)"
        ] = None

        df_report[
            "Average Heart Rate (BPM)"
        ] = None

        df_report[
            "Temperature (F)"
        ] = None

        df_report[
            "Temperature (C)"
        ] = None

        df_report[
            "Humidity (%)"
        ] = None

        df_report[
            "Latitude"
        ] = None

        df_report[
            "Longitude"
        ] = None

        df_report[
            "Min Ground Elevation (m)"
        ] = None

        df_report[
            "Max Ground Elevation (m)"
        ] = None

        df_report[
            "Hardware"
        ] = None

        df_report[
            "Source"
        ] = None

        df_report[
            "Software Version"
        ] = None

        df_report[
            "Timezone"
        ] = None


    df_report[
        "Workout Schema"
    ] = (

        df_workouts[
            "workout_schema"
        ]

    )


    if major_version in (16, 17):

        df_report[
            "Source Database"
        ] = (
            "healthdb_secure.sqlite + healthdb.sqlite"
        )

    else:

        df_report[
            "Source Database"
        ] = (
            "healthdb_secure.sqlite"
        )


    workouts_reports[
        version
    ] = df_report


print(
    "✅ Workouts report formatted."
)


for version, df_report in workouts_reports.items():

    print()
    print("=" * 70)

    print(
        f"🏋️ WORKOUTS REPORT - {version}"
    )

    print("=" * 70)


    display(
        df_report.head(20)
    )


    print()

    print(
        f"📊 Total Workouts: "
        f"{len(df_report)}"
    )

# **Wrist Temperature**


print()
print("=" * 70)
print("🌡️ WRIST TEMPERATURE STRUCTURE VERIFICATION")
print("=" * 70)


wrist_temp_secure_tables = [
    "samples",
    "quantity_samples",
    "objects",
    "data_provenances",
    "metadata_keys",
    "metadata_values"
]


wrist_temp_healthdb_tables = [
    "sources",
    "source_devices"
]


for dataset in target_datasets:

    version = dataset[
        "version"
    ]

    secure_tables = dataset[
        "healthdb_secure_tables"
    ]

    healthdb_tables = dataset[
        "healthdb_tables"
    ]


    print()
    print(
        f"📱 {version}"
    )


    print(
        "   📂 healthdb_secure.sqlite"
    )


    for table_name in wrist_temp_secure_tables:

        if table_name in secure_tables:

            print(
                f"      {table_name:<22}: ✅"
            )

        else:

            print(
                f"      {table_name:<22}: ❌"
            )


    print(
        "   📂 healthdb.sqlite"
    )


    for table_name in wrist_temp_healthdb_tables:

        if table_name in healthdb_tables:

            print(
                f"      {table_name:<22}: ✅"
            )

        else:

            print(
                f"      {table_name:<22}: ❌"
            )


print()
print("=" * 70)


def parse_wrist_cocoa_datetime(
    value
):


    if value is None:

        return None


    try:

        if pd.isna(value):

            return None

    except TypeError:

        pass


    try:

        timestamp = float(
            value
        )

    except (
        TypeError,
        ValueError
    ):

        return None


    if timestamp > 1e12:

        timestamp = (
            timestamp
            /
            1e9
        )


    if timestamp == 0:

        return None


    return parse_cocoa_datetime(
        timestamp
    )


print(
    "✅ Wrist Temperature timestamp helper is ready."
)


def clean_wrist_text(
    value
):


    if value is None:

        return None


    try:

        if pd.isna(value):

            return None

    except TypeError:

        pass


    text = (
        str(value)
        .replace(
            "\xa0",
            " "
        )
        .strip()
    )


    try:

        fixed = (

            text
            .encode(
                "latin-1",
                errors="ignore"
            )
            .decode(
                "utf-8",
                errors="ignore"
            )

        )


        if (
            fixed
            and
            len(fixed) >= len(text) - 1
        ):

            text = fixed


    except (
        UnicodeEncodeError,
        UnicodeDecodeError
    ):

        pass


    return (
        text
        if text
        else None
    )


print(
    "✅ Wrist Temperature text-cleaning helper is ready."
)


def extract_wrist_temperature(
    secure_database_path,
    health_database_path
):


    conn = open_workouts_cross_database_readonly(
        secure_database_path,
        health_database_path
    )


    query = """

    WITH surface_temp AS (

        SELECT

            mv.object_id,

            mv.numerical_value

        FROM secure_db.metadata_values AS mv


        JOIN secure_db.metadata_keys AS mk

            ON mv.key_id =
               mk.ROWID


        WHERE

            mk.key =
            '_HKPrivateMetadataKeySkinSurfaceTemperature'

    ),


    alg_ver AS (

        SELECT

            mv.object_id,

            mv.numerical_value

        FROM secure_db.metadata_values AS mv


        JOIN secure_db.metadata_keys AS mk

            ON mv.key_id =
               mk.ROWID


        WHERE

            mk.key =
            'HKAlgorithmVersion'

    )


    SELECT

        samples.start_date
            AS start_date,


        samples.end_date
            AS end_date,


        objects.creation_date
            AS creation_date,


        quantity_samples.quantity
            AS wrist_temp_c,


        (
            quantity_samples.quantity
            * 1.8
        )
        + 32
            AS wrist_temp_f,


        src.name
            AS source,


        alg_ver.numerical_value
            AS alg_version,


        surface_temp.numerical_value
            AS surf_temp_c,


        (
            surface_temp.numerical_value
            * 1.8
        )
        + 32
            AS surf_temp_f,


        sd.name
            AS device_name,


        sd.manufacturer
            AS device_manufacturer,


        sd.model
            AS device_model,


        sd.hardware
            AS hardware_version,


        sd.software
            AS software_version


    FROM secure_db.samples AS samples


    LEFT JOIN secure_db.quantity_samples
        AS quantity_samples

        ON quantity_samples.data_id =
           samples.data_id


    LEFT JOIN secure_db.objects
        AS objects

        ON samples.data_id =
           objects.data_id


    LEFT JOIN secure_db.data_provenances
        AS data_provenances

        ON objects.provenance =
           data_provenances.ROWID


    LEFT JOIN surface_temp

        ON surface_temp.object_id =
           samples.data_id


    LEFT JOIN alg_ver

        ON alg_ver.object_id =
           samples.data_id


    LEFT JOIN health_db.sources AS src

        ON src.ROWID =
           data_provenances.source_id


    LEFT JOIN health_db.source_devices AS sd

        ON sd.ROWID =
           data_provenances.device_id


    WHERE

        samples.data_type = 70


    ORDER BY

        samples.start_date DESC

    """


    df = pd.read_sql_query(
        query,
        conn
    )


    conn.close()


    return df


print(
    "✅ Wrist Temperature extractor is ready."
)


wrist_temperature_results = {}


for dataset in target_datasets:

    version = dataset[
        "version"
    ]


    print()
    print("=" * 70)

    print(
        f"🌡️ WRIST TEMPERATURE - {version}"
    )

    print("=" * 70)


    try:

        secure_db_path = dataset[
            "healthdb_secure"
        ]

        healthdb_path = dataset[
            "healthdb"
        ]


        df_wrist_temp = extract_wrist_temperature(

            secure_db_path,

            healthdb_path

        )


        wrist_temperature_results[
            version
        ] = df_wrist_temp


        print(
            "✅ Extraction successful"
        )

        print(
            f"📊 Total Wrist Temperature Records: "
            f"{len(df_wrist_temp)}"
        )

        print(
            "📂 Source: "
            "healthdb_secure.sqlite "
            "+ healthdb.sqlite"
        )


    except Exception as e:

        print(
            f"❌ Error {version}: "
            f"{type(e).__name__}: {e}"
        )


for version, df_wrist_temp in wrist_temperature_results.items():

    print()
    print("=" * 70)

    print(
        f"🌡️ RAW WRIST TEMPERATURE - {version}"
    )

    print("=" * 70)


    display(
        df_wrist_temp.head(20)
    )


    print()

    print(
        f"📊 Total Records: "
        f"{len(df_wrist_temp)}"
    )


wrist_temperature_reports = {}


for version, df_wrist_temp in wrist_temperature_results.items():

    df_report = pd.DataFrame()


    df_report[
        "Start Timestamp"
    ] = (

        df_wrist_temp[
            "start_date"
        ]
        .apply(
            parse_wrist_cocoa_datetime
        )

    )


    df_report[
        "End Timestamp"
    ] = (

        df_wrist_temp[
            "end_date"
        ]
        .apply(
            parse_wrist_cocoa_datetime
        )

    )


    df_report[
        "Date Added"
    ] = (

        df_wrist_temp[
            "creation_date"
        ]
        .apply(
            parse_wrist_cocoa_datetime
        )

    )


    df_report[
        "Wrist Temperature (C)"
    ] = (

        pd.to_numeric(
            df_wrist_temp[
                "wrist_temp_c"
            ],
            errors="coerce"
        )

    )


    df_report[
        "Wrist Temperature (F)"
    ] = (

        pd.to_numeric(
            df_wrist_temp[
                "wrist_temp_f"
            ],
            errors="coerce"
        )

    )


    df_report[
        "Algorithm Version"
    ] = (

        df_wrist_temp[
            "alg_version"
        ]

    )


    df_report[
        "Surface Temperature (C)"
    ] = (

        pd.to_numeric(
            df_wrist_temp[
                "surf_temp_c"
            ],
            errors="coerce"
        )

    )


    df_report[
        "Surface Temperature (F)"
    ] = (

        pd.to_numeric(
            df_wrist_temp[
                "surf_temp_f"
            ],
            errors="coerce"
        )

    )


    df_report[
        "Source"
    ] = (

        df_wrist_temp[
            "source"
        ]
        .apply(
            clean_wrist_text
        )

    )


    df_report[
        "Device Name"
    ] = (

        df_wrist_temp[
            "device_name"
        ]
        .apply(
            clean_wrist_text
        )

    )


    df_report[
        "Device Manufacturer"
    ] = (

        df_wrist_temp[
            "device_manufacturer"
        ]
        .apply(
            clean_wrist_text
        )

    )


    df_report[
        "Device Model"
    ] = (

        df_wrist_temp[
            "device_model"
        ]
        .apply(
            clean_wrist_text
        )

    )


    df_report[
        "Hardware Version"
    ] = (

        df_wrist_temp[
            "hardware_version"
        ]
        .apply(
            clean_wrist_text
        )

    )


    df_report[
        "Software Version"
    ] = (

        df_wrist_temp[
            "software_version"
        ]
        .apply(
            clean_wrist_text
        )

    )


    df_report[
        "Source Database"
    ] = (
        "healthdb_secure.sqlite + healthdb.sqlite"
    )


    wrist_temperature_reports[
        version
    ] = df_report


print(
    "✅ Wrist Temperature report formatted."
)


for version, df_report in wrist_temperature_reports.items():

    print()
    print("=" * 70)

    print(
        f"🌡️ WRIST TEMPERATURE REPORT - {version}"
    )

    print("=" * 70)


    display(
        df_report.head(20)
    )


    print()

    print(
        f"📊 Total Records: "
        f"{len(df_report)}"
    )

# **Create Forensic Timeline Excel Report**


output_folder = "/content/apple_health_output"

os.makedirs(
    output_folder,
    exist_ok=True
)


def datetime_to_excel_string(value):

    if pd.isna(value):
        return None

    try:
        return value.strftime(
            "%Y-%m-%d %H:%M:%S+00:00"
        )

    except (AttributeError, ValueError):
        return str(value)


summary_data = []

# **Achievements**


all_achievement_reports = []


for version, df_report in achievement_reports.items():

    df_export = df_report.copy()


    if "Created Timestamp" in df_export.columns:

        df_export[
            "Created Timestamp"
        ] = (
            df_export[
                "Created Timestamp"
            ]
            .apply(
                datetime_to_excel_string
            )
        )


    df_export.insert(
        0,
        "iOS Version",
        version
    )


    all_achievement_reports.append(
        df_export
    )


    summary_data.append({

        "iOS Version":
            version,

        "Feature":
            "Achievements",

        "Total Events":
            len(df_report),

        "Source Database":
            "healthdb_secure.sqlite"

    })


df_achievements_excel = pd.concat(
    all_achievement_reports,
    ignore_index=True
)

# **All Watch Sleep**


all_sleep_reports = []


for version, df_report in all_watch_sleep_reports.items():

    df_export = df_report.copy()


    if "Start Timestamp" in df_export.columns:

        df_export[
            "Start Timestamp"
        ] = (
            df_export[
                "Start Timestamp"
            ]
            .apply(
                datetime_to_excel_string
            )
        )


    if "End Timestamp" in df_export.columns:

        df_export[
            "End Timestamp"
        ] = (
            df_export[
                "End Timestamp"
            ]
            .apply(
                datetime_to_excel_string
            )
        )


    df_export.insert(
        0,
        "iOS Version",
        version
    )


    all_sleep_reports.append(
        df_export
    )


    summary_data.append({

        "iOS Version":
            version,

        "Feature":
            "All Watch Sleep",

        "Total Events":
            len(df_report),

        "Source Database":
            "healthdb_secure.sqlite"

    })


df_sleep_excel = pd.concat(
    all_sleep_reports,
    ignore_index=True
)

# **Headphone Audio Levels**


all_headphone_audio_reports = []


for version, df_report in headphone_audio_reports.items():

    df_export = df_report.copy()


    if "Start Timestamp" in df_export.columns:

        df_export[
            "Start Timestamp"
        ] = (
            df_export[
                "Start Timestamp"
            ]
            .apply(
                datetime_to_excel_string
            )
        )


    if "End Timestamp" in df_export.columns:

        df_export[
            "End Timestamp"
        ] = (
            df_export[
                "End Timestamp"
            ]
            .apply(
                datetime_to_excel_string
            )
        )


    df_export.insert(
        0,
        "iOS Version",
        version
    )


    all_headphone_audio_reports.append(
        df_export
    )


    summary_data.append({

        "iOS Version":
            version,

        "Feature":
            "Headphone Audio Levels",

        "Total Events":
            len(df_report),

        "Source Database":
            "healthdb_secure.sqlite + healthdb.sqlite"

    })


df_headphone_audio_excel = pd.concat(
    all_headphone_audio_reports,
    ignore_index=True
)

# **Heart Rate**


all_heart_rate_reports = []


for version, df_report in heart_rate_reports.items():

    df_export = df_report.copy()


    if "Start Timestamp" in df_export.columns:

        df_export[
            "Start Timestamp"
        ] = (
            df_export[
                "Start Timestamp"
            ]
            .apply(
                datetime_to_excel_string
            )
        )


    if "End Timestamp" in df_export.columns:

        df_export[
            "End Timestamp"
        ] = (
            df_export[
                "End Timestamp"
            ]
            .apply(
                datetime_to_excel_string
            )
        )


    if "Added Timestamp" in df_export.columns:

        df_export[
            "Added Timestamp"
        ] = (
            df_export[
                "Added Timestamp"
            ]
            .apply(
                datetime_to_excel_string
            )
        )


    df_export.insert(
        0,
        "iOS Version",
        version
    )


    all_heart_rate_reports.append(
        df_export
    )


    summary_data.append({

        "iOS Version":
            version,

        "Feature":
            "Heart Rate",

        "Total Events":
            len(df_report),

        "Source Database":
            "healthdb_secure.sqlite + healthdb.sqlite"

    })


df_heart_rate_excel = pd.concat(
    all_heart_rate_reports,
    ignore_index=True
)

# **Height**


all_height_reports = []


for version, df_report in height_reports.items():

    df_export = df_report.copy()


    if "Timestamp" in df_export.columns:

        df_export[
            "Timestamp"
        ] = (
            df_export[
                "Timestamp"
            ]
            .apply(
                datetime_to_excel_string
            )
        )


    df_export.insert(
        0,
        "iOS Version",
        version
    )


    all_height_reports.append(
        df_export
    )


    summary_data.append({

        "iOS Version":
            version,

        "Feature":
            "Height",

        "Total Events":
            len(df_report),

        "Source Database":
            "healthdb_secure.sqlite"

    })


df_height_excel = pd.concat(
    all_height_reports,
    ignore_index=True
)

# **Resting Heart Rate**


all_resting_hr_reports = []


for version, df_report in resting_hr_reports.items():


    df_export = df_report.copy()


    for col in [

        "Start Timestamp",

        "End Timestamp",

        "Added Timestamp"

    ]:


        if col in df_export.columns:


            df_export[col] = (

                df_export[col]
                .apply(
                    datetime_to_excel_string
                )

            )


    df_export.insert(
        0,
        "iOS Version",
        version
    )


    all_resting_hr_reports.append(
        df_export
    )


    summary_data.append({

        "iOS Version":
            version,


        "Feature":
            "Resting Heart Rate",


        "Total Events":
            len(df_report),


        "Source Database":
            "healthdb_secure.sqlite + healthdb.sqlite"

    })


df_resting_hr_excel = pd.concat(

    all_resting_hr_reports,

    ignore_index=True

)

# **Source Devices**


all_source_devices_reports = []


for version, df_report in source_devices_reports.items():


    df_export = df_report.copy()


    if "Creation Timestamp" in df_export.columns:


        df_export[
            "Creation Timestamp"
        ] = (

            df_export[
                "Creation Timestamp"
            ]
            .apply(
                datetime_to_excel_string
            )

        )


    df_export.insert(

        0,

        "iOS Version",

        version

    )


    all_source_devices_reports.append(
        df_export
    )


    summary_data.append({

        "iOS Version":

            version,


        "Feature":

            "Source Devices",


        "Total Events":

            len(df_report),


        "Source Database":

            "healthdb.sqlite"

    })


df_source_devices_excel = pd.concat(

    all_source_devices_reports,

    ignore_index=True

)

# **Steps**


all_steps_reports = []


for version, df_report in steps_reports.items():

    df_export = df_report.copy()


    if "Start Timestamp" in df_export.columns:

        df_export[
            "Start Timestamp"
        ] = (

            df_export[
                "Start Timestamp"
            ]
            .apply(
                datetime_to_excel_string
            )

        )


    if "End Timestamp" in df_export.columns:

        df_export[
            "End Timestamp"
        ] = (

            df_export[
                "End Timestamp"
            ]
            .apply(
                datetime_to_excel_string
            )

        )


    df_export.insert(
        0,
        "iOS Version",
        version
    )


    all_steps_reports.append(
        df_export
    )


    summary_data.append({

        "iOS Version":
            version,

        "Feature":
            "Steps",

        "Total Events":
            len(df_report),

        "Source Database":
            "healthdb_secure.sqlite"

    })


df_steps_excel = pd.concat(
    all_steps_reports,
    ignore_index=True
)


print(
    "✅ Steps data is ready for Excel."
)

# **Watch by Sleep Period**


all_watch_sleep_reports = []


for version, df_report in watch_sleep_reports.items():

    df_export = df_report.copy()


    if "Start Timestamp" in df_export.columns:

        df_export[
            "Start Timestamp"
        ] = (

            df_export[
                "Start Timestamp"
            ]
            .apply(
                datetime_to_excel_string
            )

        )


    if "End Timestamp" in df_export.columns:

        df_export[
            "End Timestamp"
        ] = (

            df_export[
                "End Timestamp"
            ]
            .apply(
                datetime_to_excel_string
            )

        )


    df_export.insert(
        0,
        "iOS Version",
        version
    )


    all_watch_sleep_reports.append(
        df_export
    )


    summary_data.append({

        "iOS Version":
            version,

        "Feature":
            "Watch By Sleep Period",

        "Total Events":
            len(df_report),

        "Source Database":
            "healthdb_secure.sqlite"

    })


df_watch_sleep_excel = pd.concat(
    all_watch_sleep_reports,
    ignore_index=True
)


print(
    "✅ Watch By Sleep Period data is ready for Excel."
)

# **Watch Worn Data**


all_watch_worn_reports = []


for version, df_report in watch_worn_reports.items():

    df_export = df_report.copy()


    if "Start Time" in df_export.columns:

        df_export[
            "Start Time"
        ] = (

            df_export[
                "Start Time"
            ]
            .apply(
                datetime_to_excel_string
            )

        )


    if "Last Worn Time" in df_export.columns:

        df_export[
            "Last Worn Time"
        ] = (

            df_export[
                "Last Worn Time"
            ]
            .apply(
                datetime_to_excel_string
            )

        )


    df_export.insert(
        0,
        "iOS Version",
        version
    )


    all_watch_worn_reports.append(
        df_export
    )


    summary_data.append({

        "iOS Version":
            version,

        "Feature":
            "Watch Worn Data",

        "Total Events":
            len(df_report),

        "Source Database":
            "healthdb_secure.sqlite"

    })


df_watch_worn_excel = pd.concat(
    all_watch_worn_reports,
    ignore_index=True
)


print(
    "✅ Watch Worn Data is ready for Excel."
)

# **Weight**


all_weight_reports = []


for version, df_report in weight_reports.items():

    df_export = df_report.copy()


    if "Timestamp" in df_export.columns:

        df_export[
            "Timestamp"
        ] = (

            df_export[
                "Timestamp"
            ]
            .apply(
                datetime_to_excel_string
            )

        )


    df_export.insert(
        0,
        "iOS Version",
        version
    )


    all_weight_reports.append(
        df_export
    )


    summary_data.append({

        "iOS Version":
            version,

        "Feature":
            "Weight",

        "Total Events":
            len(df_report),

        "Source Database":
            "healthdb_secure.sqlite"

    })


df_weight_excel = pd.concat(
    all_weight_reports,
    ignore_index=True
)


print(
    "✅ Weight data is ready for Excel."
)

# **Workouts**


all_workouts_reports = []


for version, df_report in workouts_reports.items():

    df_export = df_report.copy()


    timestamp_columns = [

        "Start Timestamp",

        "End Timestamp",

        "Added Timestamp"

    ]


    for column_name in timestamp_columns:

        if column_name in df_export.columns:

            df_export[
                column_name
            ] = (

                df_export[
                    column_name
                ]
                .apply(
                    datetime_to_excel_string
                )

            )


    df_export.insert(
        0,
        "iOS Version",
        version
    )


    all_workouts_reports.append(
        df_export
    )


    major_version = get_ios_major_version(
        version
    )


    if major_version in (16, 17):

        source_database = (
            "healthdb_secure.sqlite + healthdb.sqlite"
        )

    else:

        source_database = (
            "healthdb_secure.sqlite"
        )


    summary_data.append({

        "iOS Version":
            version,

        "Feature":
            "Workouts",

        "Total Events":
            len(df_report),

        "Source Database":
            source_database

    })


df_workouts_excel = pd.concat(
    all_workouts_reports,
    ignore_index=True
)


print(
    "✅ Workouts data is ready for Excel."
)

# **Wrist Temperature**


all_wrist_temperature_reports = []


for version, df_report in wrist_temperature_reports.items():

    df_export = df_report.copy()


    timestamp_columns = [

        "Start Timestamp",

        "End Timestamp",

        "Date Added"

    ]


    for column_name in timestamp_columns:

        if column_name in df_export.columns:

            df_export[
                column_name
            ] = (

                df_export[
                    column_name
                ]
                .apply(
                    datetime_to_excel_string
                )

            )


    df_export.insert(
        0,
        "iOS Version",
        version
    )


    all_wrist_temperature_reports.append(
        df_export
    )


    summary_data.append({

        "iOS Version":
            version,

        "Feature":
            "Wrist Temperature",

        "Total Events":
            len(df_report),

        "Source Database":
            "healthdb_secure.sqlite + healthdb.sqlite"

    })


df_wrist_temperature_excel = pd.concat(

    all_wrist_temperature_reports,

    ignore_index=True

)


print(
    "✅ Wrist Temperature data is ready for Excel."
)


df_summary = pd.DataFrame(
    summary_data
)


total_achievement_events = len(
    df_achievements_excel
)

total_sleep_events = len(
    df_sleep_excel
)

total_headphone_events = len(
    df_headphone_audio_excel
)

total_heart_rate_events = len(
    df_heart_rate_excel
)

total_height_events = len(
    df_height_excel
)

total_resting_hr_events = len(
    df_resting_hr_excel
)

total_source_devices_events = len(
    df_source_devices_excel
)

total_steps_events = len(
    df_steps_excel
)

total_watch_sleep_events = len(
    df_watch_sleep_excel
)

total_watch_worn_events = len(
    df_watch_worn_excel
)

total_weight_events = len(
    df_weight_excel
)

total_workouts_events = len(
    df_workouts_excel
)

total_wrist_temperature_events = len(
    df_wrist_temperature_excel
)


total_all_events = (

    total_achievement_events
    +
    total_sleep_events
    +
    total_headphone_events
    +
    total_heart_rate_events
    +
    total_height_events
    +
    total_resting_hr_events
    +
    total_source_devices_events
    +
    total_steps_events
    +
    total_watch_sleep_events
    +
    total_watch_worn_events
    +
    total_weight_events
    +
    total_workouts_events
    +
    total_wrist_temperature_events

)


if ios_selection == "All Versions (All Datasets)":

    output_filename = (
        "Apple_Health_Forensic_Report_All_iOS.xlsx"
    )

else:

    ios_filename = (
        ios_selection
        .replace(" ", "_")
        .replace(".", "_")
    )


    output_filename = (
        f"Apple_Health_Forensic_Report_{ios_filename}.xlsx"
    )


output_path = os.path.join(
    output_folder,
    output_filename
)


with pd.ExcelWriter(
    output_path,
    engine="openpyxl"
) as writer:


    df_summary.to_excel(
        writer,
        sheet_name="Summary",
        index=False
    )


    df_achievements_excel.to_excel(
        writer,
        sheet_name="Achievements",
        index=False
    )


    df_sleep_excel.to_excel(
        writer,
        sheet_name="All Watch Sleep",
        index=False
    )


    df_headphone_audio_excel.to_excel(
        writer,
        sheet_name="Headphone Audio",
        index=False
    )


    df_heart_rate_excel.to_excel(
        writer,
        sheet_name="Heart Rate",
        index=False
    )


    df_height_excel.to_excel(
        writer,
        sheet_name="Height",
        index=False
    )


    df_resting_hr_excel.to_excel(
        writer,
        sheet_name="Resting Heart Rate",
        index=False
    )


    df_source_devices_excel.to_excel(
        writer,
        sheet_name="Source Devices",
        index=False
    )


    df_steps_excel.to_excel(
        writer,
        sheet_name="Steps",
        index=False
    )


    df_watch_sleep_excel.to_excel(
        writer,
        sheet_name="Watch Sleep Period",
        index=False
    )


    df_watch_worn_excel.to_excel(
        writer,
        sheet_name="Watch Worn Data",
        index=False
    )


    df_weight_excel.to_excel(
        writer,
        sheet_name="Weight",
        index=False
    )


    df_workouts_excel.to_excel(
        writer,
        sheet_name="Workouts",
        index=False
    )


    df_wrist_temperature_excel.to_excel(
        writer,
        sheet_name="Wrist Temperature",
        index=False
    )


print()
print("=" * 70)
print("📄 APPLE HEALTH REPORT CREATED")
print("=" * 70)

print(
    f"📁 File : {output_filename}"
)

print(
    f"📍 Path : {output_path}"
)

print()

print(
    f"🏆 Achievements Events       : "
    f"{total_achievement_events}"
)

print(
    f"😴 All Watch Sleep Events    : "
    f"{total_sleep_events}"
)

print(
    f"🎧 Headphone Audio Events    : "
    f"{total_headphone_events}"
)

print(
    f"❤️ Heart Rate Events         : "
    f"{total_heart_rate_events}"
)

print(
    f"📏 Height Events             : "
    f"{total_height_events}"
)

print(
    f"❤️‍🩹 Resting HR Events         : "
    f"{total_resting_hr_events}"
)

print(
    f"📱 Source Devices Events     : "
    f"{total_source_devices_events}"
)

print(
    f"🚶 Steps Events              : "
    f"{total_steps_events}"
)

print(
    f"⌚ Watch Sleep Period Events : "
    f"{total_watch_sleep_events}"
)

print(
    f"⌚ Watch Worn Data Events    : "
    f"{total_watch_worn_events}"
)

print(
    f"⚖️ Weight Events             : "
    f"{total_weight_events}"
)

print(
    f"🏋️ Workouts Events           : "
    f"{total_workouts_events}"
)

print(
    f"🌡️ Wrist Temperature Events   : "
    f"{total_wrist_temperature_events}"
)

print("-" * 70)

print(
    f"📊 TOTAL EVENTS              : "
    f"{total_all_events}"
)

print("=" * 70)


from openpyxl import load_workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side
)


workbook = load_workbook(
    output_path
)


header_fill = PatternFill(
    fill_type="solid",
    fgColor="1F4E78"
)

header_font = Font(
    color="FFFFFF",
    bold=True
)

thin_border = Border(
    bottom=Side(
        style="thin",
        color="D9D9D9"
    )
)


for worksheet in workbook.worksheets:

    worksheet.freeze_panes = "A2"

    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )


    for cell in worksheet[1]:

        cell.fill = header_fill
        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


    for row in worksheet.iter_rows(
        min_row=2
    ):

        for cell in row:

            cell.alignment = Alignment(
                vertical="top"
            )

            cell.border = thin_border


    for column_cells in worksheet.columns:

        max_length = 0

        column_letter = (
            column_cells[0].column_letter
        )


        for cell in column_cells:

            try:

                cell_length = len(
                    str(cell.value)
                )

                if cell_length > max_length:
                    max_length = cell_length

            except:
                pass


        adjusted_width = min(
            max_length + 2,
            45
        )

        worksheet.column_dimensions[
            column_letter
        ].width = adjusted_width


workbook.save(
    output_path
)


print(
    "✅ Excel formatting completed."
)


from google.colab import files


print(
    f"⬇️ Downloading: {output_filename}"
)


files.download(
    output_path
)
